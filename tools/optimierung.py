"""Stangenoptimierung für die Kalkulationsmappe.

Liest die Positionen aus dem Blatt 'Auftrag', teilt die Teile je Profil + Werkstoff
+ Stangenlänge auf möglichst wenige Stangen auf und schreibt das Ergebnis in die
Blätter 'Schnittplan', 'Stangenbelegung' und 'Auswertung'.

Aufruf:
    python optimierung.py MAPPE.xlsx

Regeln (siehe Blatt 'Anleitung', A15/A16):
- Belegung = Endbeschnitt + Teile + Trennungen + Schnittverlust des letzten Schnitts;
  frei bleiben muss der Greiferrest.
- Trennung zwischen zwei Teilen: beide zugewandten Enden gerade = 1 gemeinsamer Schnitt
  (1 × Schnittverlust), sonst 2 Schnitte (2 × Schnittverlust + Abstand getrennte Schnitte).
- Teile dürfen gedreht werden. Keine Restlagerverwaltung: Reste sind Verschnitt.
"""
import argparse
import math
import shutil
import sys
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

M0, M1 = 5, 304
O0, O1 = 13, 512
S0, S1 = 5, 1004
D0, D1 = 5, 10004
EPS = 1e-6
EXAKT_MAX_TEILE = 60     # exakte Suche nur für Gruppen bis zu so vielen Teilen
EXAKT_MAX_KNOTEN = 20000  # Suchschritte je Gruppe, danach gilt das beste bis dahin gefundene Ergebnis


@dataclass(frozen=True)
class Teil:
    zeile: int          # Zeile im Blatt Auftrag
    pos: object
    artikel: str
    nr: int             # laufende Nummer innerhalb der Position (1..Stückzahl)
    laenge: float
    a_gerade: bool      # Anfang (Ende A) gerade
    b_gerade: bool
    a_schnitt: float    # Schnittlänge Ende A [mm] (für Einsparung)
    b_schnitt: float


@dataclass
class Gruppe:
    profil: str
    werkstoff: str
    stange: float
    endbeschnitt: float
    greiferrest: float
    kerf: float
    abstand: float
    teile: list = field(default_factory=list)
    bewiesen: bool = False   # Stangenzahl ist nachweislich minimal

    @property
    def nutzbar_bis(self):
        return self.stange - self.greiferrest


# ------------------------------------------------------------------ Belegung einer Stange
def trennung(kerf, abstand, ende_gerade, anfang_gerade):
    """Materialverbrauch zwischen zwei Teilen."""
    return kerf if (ende_gerade and anfang_gerade) else 2 * kerf + abstand


def lage(teil, gedreht):
    """(Anfang gerade, Ende gerade, Schnitt Anfang, Schnitt Ende) in Lage auf der Stange."""
    if gedreht:
        return teil.b_gerade, teil.a_gerade, teil.b_schnitt, teil.a_schnitt
    return teil.a_gerade, teil.b_gerade, teil.a_schnitt, teil.b_schnitt


def anordnen(teile):
    """Reihenfolge und Lage so wählen, dass möglichst viele gerade Enden aneinanderliegen.
    Schräge Enden kommen an Stangenanfang/-ende, übrige schräge Enden werden paarweise
    gegeneinander gelegt. Ergebnis: Liste (Teil, gedreht)."""
    ss = [t for t in teile if t.a_gerade and t.b_gerade]
    sn = [t for t in teile if t.a_gerade != t.b_gerade]
    nn = [t for t in teile if not t.a_gerade and not t.b_gerade]
    # Lage so, dass das schräge Ende vorn (N_vorn) bzw. hinten (N_hinten) liegt
    n_vorn = lambda t: (t, t.a_gerade)       # gedreht, wenn A gerade ist -> B (schräg) vorn
    n_hinten = lambda t: (t, not t.a_gerade)  # gedreht, wenn A schräg ist -> A hinten
    folge, rest_sn, rest_nn = [], list(sn), list(nn)
    start = n_vorn(rest_sn.pop(0)) if rest_sn else ((rest_nn.pop(0), False) if rest_nn else None)
    ende = n_hinten(rest_sn.pop(0)) if rest_sn else ((rest_nn.pop(0), False) if rest_nn else None)
    if start:
        folge.append(start)
    folge += [(t, False) for t in ss]
    # restliche: S-N | N-N ... | N-S | S-N | N-S ...
    kette, i = [], 0
    while rest_sn:
        t = rest_sn.pop(0)
        if i % 2 == 0:
            kette.append(n_hinten(t))
            while rest_nn:
                kette.append((rest_nn.pop(0), False))
        else:
            kette.append(n_vorn(t))
        i += 1
    kette += [(t, False) for t in rest_nn]
    folge += kette
    if ende:
        folge.append(ende)
    return folge


def verbrauch(g, folge):
    """Position hinter dem letzten Schnitt (inkl. Schnittverlust) ab Stangenanfang."""
    pos = g.endbeschnitt
    vorher = None
    for t, gedreht in folge:
        a_ger, b_ger, _, _ = lage(t, gedreht)
        if vorher is not None:
            pos += trennung(g.kerf, g.abstand, vorher, a_ger)
        pos += t.laenge
        vorher = b_ger
    return pos + (g.kerf if folge else 0)


def passt(g, teile):
    return verbrauch(g, anordnen(teile)) <= g.nutzbar_bis + EPS


# ------------------------------------------------------------------ Packverfahren
def sortiert(teile):
    return sorted(teile, key=lambda t: (-t.laenge, t.zeile, t.nr))


def first_fit(g, teile):
    stangen = []
    for t in sortiert(teile):
        for s in stangen:
            if passt(g, s + [t]):
                s.append(t)
                break
        else:
            stangen.append([t])
    return stangen


def best_fit(g, teile):
    stangen = []
    for t in sortiert(teile):
        beste, rest_min = None, None
        for s in stangen:
            v = verbrauch(g, anordnen(s + [t]))
            if v <= g.nutzbar_bis + EPS and (rest_min is None or g.nutzbar_bis - v < rest_min):
                beste, rest_min = s, g.nutzbar_bis - v
        if beste is None:
            stangen.append([t])
        else:
            beste.append(t)
    return stangen


def aufloesen(g, stangen):
    """Schwach belegte Stangen auflösen, wenn alle ihre Teile in andere Stangen passen."""
    stangen = [list(s) for s in stangen]
    geaendert = True
    while geaendert and len(stangen) > 1:
        geaendert = False
        for s in sorted(stangen, key=lambda s: sum(t.laenge for t in s)):
            andere = [list(x) for x in stangen if x is not s]
            ok = True
            for t in sortiert(s):
                ziel = None
                for x in sorted(andere, key=lambda x: -verbrauch(g, anordnen(x))):
                    if passt(g, x + [t]):
                        ziel = x
                        break
                if ziel is None:
                    ok = False
                    break
                ziel.append(t)
            if ok:
                stangen = andere
                geaendert = True
                break
    return stangen


def gemeinsame_schnitte(folge):
    n = 0
    for (t1, d1), (t2, d2) in zip(folge, folge[1:]):
        if lage(t1, d1)[1] and lage(t2, d2)[0]:
            n += 1
    return n


def gleichartig(a, b):
    return a.laenge == b.laenge and a.a_gerade == b.a_gerade and a.b_gerade == b.b_gerade


def exakt(g, ziel):
    """Exakte Suche (Branch and Bound) nach einer Einteilung mit weniger als `ziel` Stangen.
    Teile in der Reihenfolge von sortiert(), jedes Teil der Reihe nach in jede vorhandene Stange,
    dann in eine neue. Abgeschnitten wird, wenn selbst bei voller Nutzung aller freien Längen
    nicht weniger Stangen als bisher bestes Ergebnis möglich sind. Ergebnis: (Stangen oder None,
    Suche vollständig)."""
    teile = sortiert(g.teile)
    n = len(teile)
    nutz = g.nutzbar_bis - g.endbeschnitt
    bedarf = [t.laenge + g.kerf for t in teile]
    rest = [0.0] * (n + 1)
    for i in range(n - 1, -1, -1):
        rest[i] = rest[i + 1] + bedarf[i]
    ug = untergrenze(g)
    stangen, belegt, zu = [], [], [0] * n
    zustand = dict(beste=None, beste_n=ziel, knoten=0, abbruch=False)

    def suche(i):
        if i == n:
            zustand["beste"] = [list(s) for s in stangen]
            zustand["beste_n"] = len(stangen)
            return
        zustand["knoten"] += 1
        if zustand["knoten"] > EXAKT_MAX_KNOTEN:
            zustand["abbruch"] = True
            return
        frei = 0.0
        for b in belegt:
            frei += nutz - b
        neu = 0
        if rest[i] > frei + EPS:
            neu = math.ceil((rest[i] - frei) / nutz - EPS)
        if len(stangen) + neu >= zustand["beste_n"]:
            return
        t = teile[i]
        start = zu[i - 1] if i > 0 and gleichartig(teile[i - 1], t) else 0
        for s in range(start, len(stangen)):
            if belegt[s] + bedarf[i] > nutz + EPS or not passt(g, stangen[s] + [t]):
                continue
            stangen[s].append(t)
            belegt[s] += bedarf[i]
            zu[i] = s
            suche(i + 1)
            stangen[s].pop()
            belegt[s] -= bedarf[i]
            if zustand["abbruch"] or zustand["beste_n"] <= ug:
                return
        if len(stangen) + 1 < zustand["beste_n"]:
            stangen.append([t])
            belegt.append(bedarf[i])
            zu[i] = len(stangen) - 1
            suche(i + 1)
            stangen.pop()
            belegt.pop()

    suche(0)
    return zustand["beste"], not zustand["abbruch"] or zustand["beste_n"] <= ug


def optimiere(g):
    """Ergebnis: Liste von Stangen, je Stange die angeordnete Folge [(Teil, gedreht)].
    Setzt g.bewiesen, wenn die Stangenzahl nachweislich minimal ist."""
    kandidaten = []
    for verfahren in (first_fit, best_fit):
        st = aufloesen(g, verfahren(g, g.teile))
        folgen = [anordnen(s) for s in st]
        kandidaten.append((len(folgen), -sum(gemeinsame_schnitte(f) for f in folgen), folgen))
    kandidaten.sort(key=lambda k: (k[0], k[1]))
    folgen = kandidaten[0][2]
    ug = untergrenze(g)
    g.bewiesen = len(folgen) <= ug
    if not g.bewiesen and len(g.teile) <= EXAKT_MAX_TEILE:
        st, g.bewiesen = exakt(g, len(folgen))
        if st is not None:
            folgen = [anordnen(s) for s in st]
    # Stangen mit der höchsten Belegung zuerst (Reststange zuletzt)
    folgen.sort(key=lambda f: (-verbrauch(g, f), f[0][0].zeile))
    return folgen


def bewertung(g, n_stangen):
    """Text für die Auswertung: wie nah das Ergebnis am Optimum liegt."""
    if n_stangen <= untergrenze(g):
        return "optimal (Untergrenze erreicht)"
    if g.bewiesen:
        return "optimal (alle Einteilungen geprüft)"
    return "nicht bewiesen (Suche begrenzt)"


def manuell(g):
    """Einfache manuelle Einteilung: Eingabereihenfolge, keine Drehung, Stange voll, dann nächste."""
    stangen, aktuell = [], []
    for t in sorted(g.teile, key=lambda t: (t.zeile, t.nr)):
        if aktuell and verbrauch(g, [(x, False) for x in aktuell + [t]]) > g.nutzbar_bis + EPS:
            stangen.append(aktuell)
            aktuell = []
        aktuell.append(t)
    if aktuell:
        stangen.append(aktuell)
    return stangen


def untergrenze(g):
    nutz = g.nutzbar_bis - g.endbeschnitt
    return math.ceil(sum(t.laenge + g.kerf for t in g.teile) / nutz - EPS) if g.teile else 0


# ------------------------------------------------------------------ Mappe lesen
def _zahl(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def umfang(form, a, b):
    if form == "Rundrohr":
        return math.pi * a
    if form == "Quadratrohr":
        return 4 * a
    if form == "Rechteckrohr":
        return 2 * (a + (b or a))
    return 0.0


def lies_auftrag(wb):
    wa, wp, wm, ww = wb["Auftrag"], wb["Profile"], wb["Maschinen"], wb["Werkstoffe"]
    profile = {}
    for r in range(M0, M1 + 1):
        pid = wp.cell(r, 1).value
        if pid not in (None, ""):
            profile[pid] = dict(form=wp.cell(r, 2).value, a=_zahl(wp.cell(r, 3).value) or 0,
                                b=_zahl(wp.cell(r, 4).value), end=_zahl(wp.cell(r, 7).value),
                                greif=_zahl(wp.cell(r, 8).value))
    werkstoffe = {ww.cell(r, 1).value for r in range(M0, M1 + 1) if ww.cell(r, 1).value not in (None, "")}
    maschine = wa["C8"].value
    mrow = next((r for r in range(M0, M1 + 1) if maschine and wm.cell(r, 1).value == maschine), None)
    if mrow is None:
        raise ValueError(f"Maschine '{maschine}' im Auftragskopf (C8) nicht im Blatt 'Maschinen' gefunden")
    m = dict(end=_zahl(wm.cell(mrow, 3).value) or 0, greif=_zahl(wm.cell(mrow, 4).value) or 0,
             kerf=_zahl(wm.cell(mrow, 5).value) or 0, maxlen=_zahl(wm.cell(mrow, 6).value),
             abstand=_zahl(wm.cell(mrow, 10).value) or 0)
    std_stange = _zahl(wa["C9"].value)

    gruppen, fehler = {}, []
    for r in range(O0, O1 + 1):
        pos, art, pid, wst = (wa.cell(r, c).value for c in (1, 2, 4, 5))
        laenge, stk = _zahl(wa.cell(r, 6).value), wa.cell(r, 7).value
        if all(v in (None, "") for v in (pid, wst, laenge, stk)):
            continue
        name = art or f"Zeile {r}"
        stk_n = int(stk) if isinstance(stk, (int, float)) and stk >= 1 and int(stk) == stk else 0

        def fehl(grund):
            fehler.append(dict(zeile=r, pos=pos, artikel=name, anzahl=stk_n, grund=grund))
        if pid not in profile:
            fehl("Profil unbekannt"); continue
        if wst not in werkstoffe:
            fehl("Werkstoff unbekannt"); continue
        if not laenge or laenge <= 0:
            fehl("Teilelänge fehlt"); continue
        if not stk_n:
            fehl("Stückzahl fehlt oder ungültig"); continue
        stange = _zahl(wa.cell(r, 10).value) or std_stange
        if not stange:
            fehl("Stangenlänge fehlt"); continue
        if m["maxlen"] and stange > m["maxlen"]:
            fehl("Stange länger als Maschine zulässt"); continue
        p = profile[pid]
        end = p["end"] if p["end"] is not None else m["end"]
        greif = p["greif"] if p["greif"] is not None else m["greif"]
        if end + laenge + m["kerf"] > stange - greif + EPS:
            fehl("Teil länger als Nutzlänge der Stange"); continue
        key = (pid, wst, stange)
        g = gruppen.setdefault(key, Gruppe(pid, wst, stange, end, greif, m["kerf"], m["abstand"]))
        a_ger = wa.cell(r, 25).value != "nein"   # leer = gerade
        b_ger = wa.cell(r, 27).value != "nein"
        u = umfang(p["form"], p["a"], p["b"])
        h = _zahl(wa.cell(r, 8).value)
        if h:   # Schneidweg enthält beide Endschnitte -> Ersatzwert nicht größer als die Hälfte
            u = min(u, h / 2)
        a_s = _zahl(wa.cell(r, 26).value) or u
        b_s = _zahl(wa.cell(r, 28).value) or u
        for i in range(1, stk_n + 1):
            g.teile.append(Teil(r, pos, name, i, laenge, a_ger, b_ger, a_s, b_s))
    return list(gruppen.values()), fehler


# ------------------------------------------------------------------ Ergebnis schreiben
def leeren(ws, r0, r1, spalten):
    for r in range(r0, r1 + 1):
        for c in spalten:
            ws.cell(r, c).value = None


def schreibe(wb, gruppen, ergebnisse, fehler):
    wsp, wdt, wau = wb["Schnittplan"], wb["Stangenbelegung"], wb["Auswertung"]
    leeren(wsp, S0, S1, (1, 2, 3, 4, 5, 6, 9, 14, 15, 16))
    for row in wdt.iter_rows(min_row=D0, max_row=max(wdt.max_row, D0)):
        for c in row:
            c.value = None
    for row in wau.iter_rows(min_row=4, max_row=max(wau.max_row, 4)):
        for c in row:
            c.value = None
            c.fill = PatternFill()
            c.font = Font(name="Arial")
            c.border = Border()

    f_ar = Font(name="Arial")
    nr, sr, dr = 0, S0, D0
    for g, folgen in zip(gruppen, ergebnisse):
        for folge in folgen:
            nr += 1
            if sr > S1:
                raise RuntimeError("Mehr als 1000 Stangen: Blatt 'Schnittplan' erweitern")
            pos = g.endbeschnitt
            vorher, verlust, belegung, n_gem = None, 0.0, [], 0
            for platz, (t, gedreht) in enumerate(folge, start=1):
                a_ger, b_ger, a_s, b_s = lage(t, gedreht)
                eingespart_mm, eingespart_e = 0.0, 0
                if vorher is None:
                    art = "Stangenanfang"
                else:
                    gap = trennung(g.kerf, g.abstand, vorher[0], a_ger)
                    gemeinsam = vorher[0] and a_ger
                    art = "gemeinsamer Schnitt" if gemeinsam else "2 Schnitte"
                    if gemeinsam:
                        n_gem += 1
                        eingespart_mm, eingespart_e = min(vorher[1], a_s), 1
                    pos += gap
                    verlust += gap
                von, bis = pos, pos + t.laenge
                if dr > D1:
                    raise RuntimeError("Mehr als 10000 Teile: Blatt 'Stangenbelegung' erweitern")
                werte = [nr, platz, t.pos, t.artikel, t.zeile, t.laenge, round(von, 2), round(bis, 2),
                         "gedreht" if gedreht else "normal", art, round(eingespart_mm, 1) or None,
                         eingespart_e or None]
                for c, v in enumerate(werte, start=1):
                    cell = wdt.cell(dr, c, v)
                    cell.font = f_ar
                dr += 1
                pos = bis
                vorher = (b_ger, b_s)
                belegung.append(f"{t.pos if t.pos is not None else t.artikel}×{t.laenge:g}")
            verlust += g.kerf
            werte = {1: nr, 2: g.profil, 3: g.werkstoff, 4: g.stange, 5: g.endbeschnitt, 6: g.greiferrest,
                     9: round(verlust, 2), 14: len(folge), 15: n_gem, 16: " | ".join(belegung)}
            for c, v in werte.items():
                wsp.cell(sr, c, v)
            sr += 1

    # Auswertung
    kopf = ["Profil-ID", "Werkstoff-ID", "Stangenlänge [mm]", "Teile [Stk]", "Stangen optimiert",
            "Stangen manuell", "Untergrenze Stangen", "Ersparnis Stangen", "Materialbedarf [m]",
            "Teilelänge [m]", "Verschnitt [m]", "Ausnutzung optimiert", "Ausnutzung manuell",
            "Gemeinsame Schnitte", "Prüfung Stangen", "Über Untergrenze [Stangen]", "Bewertung"]
    hfill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="BFBFBF")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, k in enumerate(kopf, start=1):
        cell = wau.cell(4, c, k)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = hfill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = bd
        wau.column_dimensions[cell.column_letter].width = 13
    wau.column_dimensions["A"].width = 20
    wau.column_dimensions["Q"].width = 34
    wau.row_dimensions[4].height = 45
    sp = lambda col: f"Schnittplan!${col}${S0}:${col}${S1}"
    r = 5
    for g, folgen in zip(gruppen, ergebnisse):
        man = manuell(g)
        crit = f'{sp("B")},$A{r},{sp("C")},$B{r},{sp("D")},$C{r}'
        laenge_man = sum(t.laenge for t in g.teile)
        werte = [g.profil, g.werkstoff, g.stange,
                 f"=SUMIFS({sp('G')},{crit})", f"=COUNTIFS({crit})", len(man), untergrenze(g), f"=F{r}-E{r}",
                 f"=E{r}*C{r}/1000", f"=SUMIFS({sp('H')},{crit})/1000", f"=I{r}-J{r}", f"=IF(I{r}=0,0,J{r}/I{r})",
                 laenge_man / (len(man) * g.stange) if man else 0, f"=SUMIFS({sp('O')},{crit})",
                 f'=IF(COUNTIFS({crit},{sp("M")},"ÜBERBELEGT")>0,"ÜBERBELEGT","OK")', f"=E{r}-G{r}",
                 bewertung(g, len(folgen))]
        for c, v in enumerate(werte, start=1):
            cell = wau.cell(r, c, v)
            cell.font = f_ar
            cell.border = bd
            cell.number_format = {3: "#,##0", 9: "#,##0.00", 10: "#,##0.00", 11: "#,##0.00", 12: "0.0%",
                                  13: "0.0%"}.get(c, "General")
        r += 1
    if gruppen:
        cell = wau.cell(r, 1, "Summe")
        cell.font = Font(name="Arial", bold=True)
        for c in (4, 5, 6, 7, 8, 9, 10, 11, 14, 16):
            col = wau.cell(5, c).column_letter
            cell = wau.cell(r, c, f"=SUM({col}5:{col}{r - 1})")
            cell.font = Font(name="Arial", bold=True)
            cell.number_format = "#,##0.00" if c in (9, 10, 11) else "#,##0"
        wau.cell(r, 12, f"=IF(I{r}=0,0,J{r}/I{r})").number_format = "0.0%"
        r += 1

    r += 1
    wau.cell(r, 1, "Nicht eingeplante Teile").font = Font(name="Arial", bold=True, size=12)
    r += 1
    for c, k in enumerate(["Zeile Auftrag", "Pos.", "Artikelnummer", "Anzahl", "Grund"], start=1):
        cell = wau.cell(r, c, k)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = hfill
    r += 1
    if not fehler:
        wau.cell(r, 1, "keine").font = f_ar
    for f in fehler:
        for c, v in enumerate([f["zeile"], f["pos"], f["artikel"], f["anzahl"], f["grund"]], start=1):
            wau.cell(r, c, v).font = f_ar
        r += 1
    if "Listen" in wb.sheetnames:
        wb["Listen"]["E5"] = "nein"   # Optimierung aktuell (siehe Auftrag C10)
    wb.calculation.fullCalcOnLoad = True


def pruefe(gruppen, ergebnisse):
    """Unabhängige Kontrolle: keine Stange überbelegt, jedes Teil genau einmal eingeplant."""
    for g, folgen in zip(gruppen, ergebnisse):
        ids = sorted((t.zeile, t.nr) for f in folgen for t, _ in f)
        soll = sorted((t.zeile, t.nr) for t in g.teile)
        if ids != soll:
            raise AssertionError(f"{g.profil}: Teile fehlen oder doppelt")
        for f in folgen:
            if verbrauch(g, f) > g.nutzbar_bis + EPS:
                raise AssertionError(f"{g.profil}: Stange überbelegt")


def optimiere_mappe(mappe, sicherung=True):
    mappe = Path(mappe)
    if sicherung:
        shutil.copy2(mappe, mappe.with_suffix(".bak" + mappe.suffix))
    wb = load_workbook(mappe, keep_vba=mappe.suffix.lower() == ".xlsm")
    gruppen, fehler = lies_auftrag(wb)
    ergebnisse = [optimiere(g) for g in gruppen]
    pruefe(gruppen, ergebnisse)
    schreibe(wb, gruppen, ergebnisse, fehler)
    wb.save(mappe)
    bericht = []
    for g, folgen in zip(gruppen, ergebnisse):
        bericht.append(dict(profil=g.profil, werkstoff=g.werkstoff, stange=g.stange, teile=len(g.teile),
                            stangen=len(folgen), manuell=len(manuell(g)), untergrenze=untergrenze(g),
                            gemeinsam=sum(gemeinsame_schnitte(f) for f in folgen)))
    return bericht, fehler


def main():
    ap = argparse.ArgumentParser(description="Stangenoptimierung für die Kalkulationsmappe")
    ap.add_argument("mappe")
    a = ap.parse_args()
    bericht, fehler = optimiere_mappe(a.mappe)
    for b in bericht:
        print(f"{b['profil']:16s} {b['werkstoff']:12s} {b['stange']:>6g} mm: {b['teile']:4d} Teile -> "
              f"{b['stangen']:3d} Stangen (manuell {b['manuell']}, Untergrenze {b['untergrenze']}), "
              f"{b['gemeinsam']} gemeinsame Schnitte")
    for f in fehler:
        print(f"NICHT EINGEPLANT Zeile {f['zeile']} {f['artikel']}: {f['grund']}")


if __name__ == "__main__":
    main()
