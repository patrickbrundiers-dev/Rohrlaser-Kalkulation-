import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

import os, sys
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "..", "mappe", "Rohrlaser_Kalkulation_v02.xlsm")
F = "Arial"
N_MASTER = 300          # Datenzeilen je Stammdatenblatt (Zeile 5..304)
M0, M1 = 5, 5 + N_MASTER - 1
O0, O1 = 13, 13 + 500 - 1   # Auftragspositionen Zeile 13..512
S0, S1 = 5, 5 + 1000 - 1    # Schnittplan: Stangen Zeile 5..1004
D0, D1 = 5, 5 + 10000 - 1   # Stangenbelegung: Teile Zeile 5..10004

f_title = Font(name=F, size=14, bold=True)
f_bold = Font(name=F, bold=True)
f_head = Font(name=F, bold=True, color="FFFFFF")
f_in = Font(name=F, color="0000FF")
f_calc = Font(name=F, color="000000")
f_link = Font(name=F, color="008000")
f_note = Font(name=F, italic=True, color="595959", size=9)
fill_head = PatternFill("solid", fgColor="1F4E78")
fill_in = PatternFill("solid", fgColor="FFF2CC")
fill_calc = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def header(ws, row, cols):
    """cols: list of (title, width, kind) kind in in/calc"""
    for i, (t, w, kind) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = f_head
        c.fill = fill_head
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = border
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 45


def style_rows(ws, r0, r1, cols, fmts=None):
    fmts = fmts or {}
    for r in range(r0, r1 + 1):
        for i, (_, _, kind) in enumerate(cols, start=1):
            c = ws.cell(row=r, column=i)
            c.border = border
            if kind == "in":
                c.fill = fill_in
                c.font = f_in
            elif kind == "link":
                c.fill = fill_calc
                c.font = f_link
            else:
                c.fill = fill_calc
                c.font = f_calc
            if i in fmts:
                c.number_format = fmts[i]


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f_title
    ws["A2"] = sub
    ws["A2"].font = f_note


def add_dv(ws, dv, rng):
    ws.add_data_validation(dv)
    dv.add(rng)


def num_dv(op_min, op_max, whole=False, title_="Ungültige Eingabe", msg=""):
    dv = DataValidation(type="whole" if whole else "decimal", operator="between",
                        formula1=str(op_min), formula2=str(op_max), allow_blank=True)
    dv.error = msg or f"Bitte einen Wert zwischen {op_min} und {op_max} eingeben."
    dv.errorTitle = title_
    dv.showErrorMessage = True
    return dv


def list_dv(name, msg):
    dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True)
    dv.error = msg
    dv.errorTitle = "Ungültige Auswahl"
    dv.showErrorMessage = True
    return dv


# ---------------------------------------------------------------- Anleitung
ws = wb.active
ws.title = "Anleitung"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 100
title(ws, "Kalkulation Rohrlaserteile – Version 0.2",
      f"Stand {datetime.date(2026, 10, 17).strftime('%d.%m.%Y')}. Optimierung und STEP-Import laufen als Makros in dieser Mappe, ohne Installation.")
rows = [
    ("Farblegende", None),
    ("Gelb / blaue Schrift", "Eingabefeld. Nur hier Werte eintragen."),
    ("Grau / schwarze Schrift", "Berechnet (Formel). Nicht überschreiben."),
    ("Grau / grüne Schrift", "Aus anderem Blatt nachgeschlagen (Formel). Nicht überschreiben."),
    ("", ""),
    ("Blätter", None),
    ("Auftrag", "Auftragskopf (Auftrag, Kunde, Maschine, Standard-Stangenlänge) und bis zu 500 Positionen. Spalte W zeigt je Zeile den Prüfstatus."),
    ("Profile", "Profil-Stammdaten: Form, Abmessungen, Wandstärke, Querschnitt; optional profilspezifischer Endbeschnitt/Greiferrest."),
    ("Werkstoffe", "Werkstoff/Materialqualität mit Dichte."),
    ("Preise", "Materialpreis je Profil + Werkstoff, wahlweise €/kg oder €/Stange."),
    ("Maschinen", "Maschinenparameter: Endbeschnitt, Greiferrest, Schnittverlust, max. Stangenlänge, Stundensatz."),
    ("Schneiddaten", "Vorschub je Konturklasse (groß/Trennschnitt, mittel, klein) und Einstechzeit je Maschine, Werkstoff und Wandstärke (für Laserzeit in Schritt 4)."),
    ("Gewinde", "Kernlochdurchmesser je Gewinde. Der STEP-Import schlägt anhand dieser Tabelle Gewinde vor."),
    ("STEP-Import", "Protokoll aller Messwerte je importiertem Teil."),
    ("Kalkulation", "Material-, Laser-, Gewinde- und Rüstkosten je Teil und Position, Zuschläge, Verkaufspreis."),
    ("Schnittplan", "Ergebnis der Optimierung: eine Zeile je Stange mit Belegung, Rest, Ausnutzung und Prüfung gegen Überbelegung."),
    ("Stangenbelegung", "Ergebnis der Optimierung: eine Zeile je Teil mit Lage auf der Stange (von/bis) und Art der Trennung."),
    ("Auswertung", "Vergleich optimierte gegen einfache manuelle Einteilung, Liste nicht eingeplanter Teile."),
    ("Listen", "Auswahllisten (Profilform, Preisbasis)."),
    ("", ""),
    ("Arbeitsablauf", None),
    ("1.", "Stammdaten pflegen (Profile, Werkstoffe, Preise, Maschinen, Schneiddaten). Die vorhandenen Zeilen sind BEISPIELWERTE und müssen durch eigene Daten ersetzt werden."),
    ("2.", "Auftragskopf ausfüllen, Maschine wählen, Standard-Stangenlänge prüfen (Vorgabe 6.000 mm)."),
    ("3.", "Positionen eintragen (Profil und Werkstoff per Auswahlliste) oder im Blatt 'Auftrag' auf 'STEP importieren' klicken und STEP-Dateien wählen. Dafür muss der Ordner 'STEP-Import' (mit STEP-Import.exe) neben der Mappe liegen."),
    ("4.", "Status-Spalte prüfen: FEHLER muss behoben werden, WARNUNG/HINWEIS sollte geprüft werden."),
    ("5.", "Im Blatt 'Auftrag' auf 'Optimieren' klicken: füllt Schnittplan, Stangenbelegung und Auswertung. Danach Blatt 'Kalkulation' prüfen. Nach jeder Änderung am Auftrag erneut optimieren."),
    ("", ""),
    ("Annahmen", None),
    ("A1", "Microsoft 365 / Excel 2019 oder neuer unter Windows, Makros müssen aktiviert sein (beim Öffnen 'Inhalt aktivieren' bzw. Datei als vertrauenswürdig markieren). Keine Installation nötig: Optimierung und Übernahme laufen in VBA, die STEP-Geometrie misst STEP-Import.exe (portable, liegt im Ordner 'STEP-Import' neben der Mappe). Die Mappe darf dabei offen bleiben."),
    ("A2", "Ein Auftrag je Blatt 'Auftrag'. Alle Positionen laufen auf der im Kopf gewählten Maschine."),
    ("A3", "Stangenlänge: Standard aus dem Auftragskopf, je Position abweichend überschreibbar (Spalte J)."),
    ("A4", "Endbeschnitt und Greiferrest kommen von der Maschine; ist am Profil ein abweichender Wert eingetragen, gilt der Profilwert."),
    ("A5", "Nutzlänge je Stange = Stangenlänge − Endbeschnitt − Greiferrest. Ein Teil passt nur, wenn Teilelänge + Schnittverlust ≤ Nutzlänge; die genaue Belegungsregel steht in A15."),
    ("A6", "Querschnitt berechnet ohne Eckradien (Quadrat-/Rechteckrohr ca. 1–4 % zu schwer). Für exakte Gewichte den Herstellerwert in 'Fläche manuell' eintragen; Sonderprofile benötigen immer 'Fläche manuell'."),
    ("A7", "Metergewicht [kg/m] = Querschnitt [mm²] × Dichte [kg/dm³] / 1000."),
    ("A8", "Keine Restlagerverwaltung: Reststücke werden nur als Verschnitt ausgewiesen (Entscheidung vom 15.10.2026)."),
    ("A9", "Kapazität: 300 Zeilen je Stammdatenblatt, 500 Auftragspositionen. Erweiterbar, indem Formelzeilen nach unten kopiert und die Namen in Formeln > Namens-Manager angepasst werden."),
    ("A10", "Einheiten: Längen in mm, Gewicht in kg, Preise in € netto."),
    ("A11", "Gemeinsamer Trennschnitt: Liegen zwei Teile auf der Stange nebeneinander und sind beide zugewandten Enden gerade (senkrecht), wird nur ein Schnitt gerechnet, sonst zwei. Dafür stehen je Teil Anfangs- und Endschnitt (gerade ja/nein, Schnittlänge) in Spalten Y–AB; Spalte H enthält immer beide Enden voll. Die Verrechnung erfolgt in der Optimierung (Schritt 3). Beim gemeinsamen Schnitt zählt auch der Schnittverlust nur einmal."),
    ("A12", "Schneidweg (Spalte H) = Anfangsschnitt + Endschnitt + Innenkonturen (AC). Gravuren (nicht durchgehende Konturen) stehen separat in AD und sind nicht im Schneidweg enthalten."),
    ("A13", "Gewinde: Der STEP-Import erkennt durchgehende Rundbohrungen (AE) und schlägt bei passendem Kernloch ein Gewinde vor (AF). Gerechnet wird nur, was nach manueller Prüfung in AG eingetragen ist (Format z. B. 4xM6; 2xM8). Solange AF gefüllt und AG leer ist, zeigt der Status einen Hinweis."),
    ("A15", "Optimierung: je Profil + Werkstoff + Stangenlänge getrennt (keine Vermischung). Belegung je Stange = Endbeschnitt + Teile + Trennungen + letzter Schnittverlust; frei bleiben muss der Greiferrest. Trennung: beide Enden gerade = 1 gemeinsamer Schnitt (1 × Schnittverlust), sonst 2 Schnitte (2 × Schnittverlust + Abstand getrennte Schnitte). Teile dürfen gedreht werden, damit gerade Enden aneinanderliegen; schräge Enden kommen bevorzugt an Stangenanfang oder -ende. Fehlt die Angabe 'gerade', gilt das Ende als gerade."),
    ("A16", "Verfahren: First-Fit- und Best-Fit-Decreasing mit anschließendem Auflösen schwach belegter Stangen; das bessere Ergebnis zählt. Untergrenze = Summe(Teilelänge + Schnittverlust) / (Stangenlänge − Endbeschnitt − Greiferrest), aufgerundet. Vergleich: manuelle Einteilung = Positionen in Eingabereihenfolge, Stange voll, dann nächste."),
    ("A17", "Kalkulation: Materialkosten je Gruppe = Summe der Stangenkosten (Stangenlänge × kg/m × €/kg oder €/Stange anteilig zur Lieferlänge). Umlage auf Positionen nach Längenanteil, damit trägt jedes Teil den Verschnitt anteilig. Laserzeit = Schneidweg je Konturklasse / Vorschub + Einstiche × Einstechzeit + Gravur / Vorschub Gravur, abzüglich der Einsparung durch gemeinsame Schnitte. Gewindezeit nur aus 'Gewinde geprüft' × Zeit je Gewinde (Blatt Gewinde). Rüstzeit und Stangenwechsel werden gleichmäßig auf alle Teile umgelegt. Fertigung mit Maschinenstundensatz."),
    ("A14", "Konturgröße: Innenkonturen werden nach ihrem größten Hüllmaß in klein / mittel / groß eingeteilt (AH–AM). Die Grenzen stehen je Maschine im Blatt 'Maschinen' (Vorgabe: klein ≤ 10 mm, mittel ≤ 50 mm, darüber groß). Anfangs- und Endschnitt werden nicht eingeteilt. Beim Import gilt die im Auftragskopf gewählte Maschine. Für die Laserzeit gilt je Klasse der eigene Vorschub aus 'Schneiddaten'; Anfangs-/Endschnitt mit Vorschub groß; fehlt mittel oder klein, wird der nächstgrößere verwendet."),
]
r = 4
for a, b in rows:
    ws.cell(row=r, column=1, value=a)
    if b is None:
        ws.cell(row=r, column=1).font = f_bold
    else:
        ws.cell(row=r, column=1).font = Font(name=F)
        c = ws.cell(row=r, column=2, value=b)
        c.font = Font(name=F)
        c.alignment = wrap
    r += 1
ws["A5"].fill = fill_in; ws["A5"].font = f_in
ws["A6"].fill = fill_calc
ws["A7"].fill = fill_calc; ws["A7"].font = f_link

# ---------------------------------------------------------------- Listen
wl = wb.create_sheet("Listen")
title(wl, "Auswahllisten", "Werte für Dropdowns. Ergänzungen hier eintragen.")
wl["A4"], wl["B4"] = "Profilform", "Preisbasis"
for c in (wl["A4"], wl["B4"]):
    c.font = f_head; c.fill = fill_head
for i, v in enumerate(["Rundrohr", "Quadratrohr", "Rechteckrohr", "Sonderprofil"], start=5):
    wl.cell(row=i, column=1, value=v).font = f_in
wl["C4"] = "Ja/Nein"
wl["C4"].font = f_head; wl["C4"].fill = fill_head
for i, v in enumerate(["ja", "nein"], start=5):
    wl.cell(row=i, column=3, value=v).font = f_in
wb.defined_names["lst_JaNein"] = DefinedName("lst_JaNein", attr_text="Listen!$C$5:$C$6")
for i, v in enumerate(["€/kg", "€/Stange"], start=5):
    wl.cell(row=i, column=2, value=v).font = f_in
wl.column_dimensions["A"].width = 18
wl.column_dimensions["B"].width = 14
wb.defined_names["lst_Profilform"] = DefinedName("lst_Profilform", attr_text="Listen!$A$5:$A$20")
wb.defined_names["lst_Preisbasis"] = DefinedName("lst_Preisbasis", attr_text="Listen!$B$5:$B$6")

# ---------------------------------------------------------------- Werkstoffe
ww = wb.create_sheet("Werkstoffe", 2)
title(ww, "Werkstoffe / Materialqualität", "Werkstoff-ID ist der Schlüssel für Preise, Schneiddaten und Auftrag. Beispielwerte ersetzen.")
cols_w = [("Werkstoff-ID", 18, "in"), ("Werkstoff", 16, "in"), ("Werkstoff-Nr.", 14, "in"),
          ("Norm / Qualität", 30, "in"), ("Dichte [kg/dm³]", 12, "in"), ("Bemerkung", 30, "in")]
header(ww, 4, cols_w)
style_rows(ww, M0, M1, cols_w, {5: "0.00"})
for i, row in enumerate([
    ("S235-10219", "S235JRH", "1.0039", "EN 10219 kaltgefertigt, geschweißt", 7.85, "Beispiel"),
    ("S355-10219", "S355J2H", "1.0576", "EN 10219 kaltgefertigt, geschweißt", 7.85, "Beispiel"),
    ("1.4301", "X5CrNi18-10", "1.4301", "EN 10217-7 / ASTM A554", 7.90, "Beispiel"),
    ("AW6060-T66", "EN AW-6060 T66", "3.3206", "EN 755-2", 2.70, "Beispiel"),
], start=M0):
    for j, v in enumerate(row, start=1):
        ww.cell(row=i, column=j, value=v)
add_dv(ww, num_dv(0.5, 25, msg="Dichte in kg/dm³ (z. B. Stahl 7,85, Alu 2,70)."), f"E{M0}:E{M1}")
dup = DataValidation(type="custom", formula1=f"=COUNTIF($A${M0}:$A${M1},A{M0})=1", allow_blank=True)
dup.error = "Diese Werkstoff-ID existiert bereits."; dup.showErrorMessage = True
add_dv(ww, dup, f"A{M0}:A{M1}")
ww.freeze_panes = "B5"
wb.defined_names["lst_Werkstoff"] = DefinedName("lst_Werkstoff", attr_text=f"Werkstoffe!$A${M0}:$A${M1}")

# ---------------------------------------------------------------- Profile
wp = wb.create_sheet("Profile", 2)
title(wp, "Profile / Abmessungen",
      "Abmessung A = Ø (Rundrohr) bzw. Breite; Abmessung B = Höhe (nur Rechteckrohr). Endbeschnitt/Greiferrest nur eintragen, wenn vom Maschinenwert abweichend.")
cols_p = [("Profil-ID", 20, "in"), ("Profilform", 15, "in"), ("Abmessung A [mm]", 12, "in"),
          ("Abmessung B [mm]", 12, "in"), ("Wandstärke [mm]", 12, "in"), ("Fläche manuell [mm²] (optional)", 14, "in"),
          ("Endbeschnitt abweichend [mm]", 14, "in"), ("Greiferrest abweichend [mm]", 14, "in"),
          ("Bemerkung", 26, "in"), ("Fläche berechnet [mm²]", 13, "calc"), ("Fläche wirksam [mm²]", 13, "calc")]
header(wp, 4, cols_p)
style_rows(wp, M0, M1, cols_p, {3: "0.0", 4: "0.0", 5: "0.00", 6: "0.0", 7: "0", 8: "0", 10: "0.0", 11: "0.0"})
for i, row in enumerate([
    ("RR 42,4x2,0", "Rundrohr", 42.4, None, 2.0, None, None, None, "Beispiel"),
    ("RR 48,3x2,6", "Rundrohr", 48.3, None, 2.6, None, None, None, "Beispiel"),
    ("QR 40x40x3", "Quadratrohr", 40, None, 3.0, None, None, None, "Beispiel"),
    ("RE 60x40x3", "Rechteckrohr", 60, 40, 3.0, None, None, None, "Beispiel"),
    ("RE 120x60x4", "Rechteckrohr", 120, 60, 4.0, None, 30, 120, "Beispiel: profilspezifischer Endbeschnitt/Greiferrest"),
], start=M0):
    for j, v in enumerate(row, start=1):
        if v is not None:
            wp.cell(row=i, column=j, value=v)
for r in range(M0, M1 + 1):
    hb = f'IF(B{r}="Quadratrohr",C{r},D{r})'
    wp[f"J{r}"] = (f'=IF(OR(A{r}="",C{r}="",E{r}=""),"",IF(B{r}="Rundrohr",PI()/4*(C{r}^2-(C{r}-2*E{r})^2),'
                   f'IF(OR(B{r}="Quadratrohr",AND(B{r}="Rechteckrohr",D{r}<>"")),C{r}*{hb}-(C{r}-2*E{r})*({hb}-2*E{r}),"")))')
    wp[f"K{r}"] = f'=IF(A{r}="","",IF(F{r}<>"",F{r},J{r}))'
add_dv(wp, list_dv("lst_Profilform", "Bitte Profilform aus der Liste wählen."), f"B{M0}:B{M1}")
add_dv(wp, num_dv(1, 1000, msg="Abmessung in mm (1–1000)."), f"C{M0}:D{M1}")
wt = DataValidation(type="custom", formula1=f"=AND(E{M0}>0,E{M0}*2<C{M0})", allow_blank=True)
wt.error = "Wandstärke muss > 0 und kleiner als die halbe Abmessung A sein."; wt.showErrorMessage = True
add_dv(wp, wt, f"E{M0}:E{M1}")
add_dv(wp, num_dv(1, 100000, msg="Querschnittsfläche in mm²."), f"F{M0}:F{M1}")
add_dv(wp, num_dv(0, 2000, msg="Wert in mm (0–2000)."), f"G{M0}:H{M1}")
dup = DataValidation(type="custom", formula1=f"=COUNTIF($A${M0}:$A${M1},A{M0})=1", allow_blank=True)
dup.error = "Diese Profil-ID existiert bereits."; dup.showErrorMessage = True
add_dv(wp, dup, f"A{M0}:A{M1}")
wp["F4"].comment = Comment("Herstellerwert eintragen, wenn exaktes Gewicht nötig (Eckradien) oder bei Sonderprofil.", "Kalkulation")
wp.freeze_panes = "B5"
wb.defined_names["lst_Profil"] = DefinedName("lst_Profil", attr_text=f"Profile!$A${M0}:$A${M1}")

# ---------------------------------------------------------------- Preise
wpr = wb.create_sheet("Preise", 4)
title(wpr, "Materialpreise", "Je Kombination Profil + Werkstoff genau eine Zeile. Preisbasis €/kg oder €/Stange (bezogen auf die Liefer-Stangenlänge).")
cols_pr = [("Profil-ID", 20, "in"), ("Werkstoff-ID", 16, "in"), ("Preisbasis", 11, "in"), ("Preis [€]", 11, "in"),
           ("Liefer-Stangenlänge [mm]", 13, "in"), ("Gültig ab", 12, "in"), ("Lieferant", 20, "in"),
           ("Bemerkung", 22, "in"), ("Prüfung", 26, "calc")]
header(wpr, 4, cols_pr)
style_rows(wpr, M0, M1, cols_pr, {4: "#,##0.00", 5: "#,##0", 6: "DD.MM.YYYY"})
for i, row in enumerate([
    ("RR 42,4x2,0", "S235-10219", "€/kg", 1.45, 6000, datetime.date(2026, 10, 1), "Beispiel-Lieferant", "Beispiel"),
    ("RR 48,3x2,6", "S235-10219", "€/kg", 1.45, 6000, datetime.date(2026, 10, 1), "Beispiel-Lieferant", "Beispiel"),
    ("QR 40x40x3", "S235-10219", "€/kg", 1.50, 6000, datetime.date(2026, 10, 1), "Beispiel-Lieferant", "Beispiel"),
    ("RE 60x40x3", "S355-10219", "€/kg", 1.60, 6000, datetime.date(2026, 10, 1), "Beispiel-Lieferant", "Beispiel"),
    ("RR 42,4x2,0", "1.4301", "€/Stange", 62.00, 6000, datetime.date(2026, 10, 1), "Beispiel-Lieferant", "Beispiel"),
], start=M0):
    for j, v in enumerate(row, start=1):
        wpr.cell(row=i, column=j, value=v)
for r in range(M0, M1 + 1):
    wpr[f"I{r}"] = (f'=IF(A{r}&B{r}="","",IF(ISNA(MATCH(A{r},lst_Profil,0)),"FEHLER: Profil unbekannt",'
                    f'IF(ISNA(MATCH(B{r},lst_Werkstoff,0)),"FEHLER: Werkstoff unbekannt",'
                    f'IF(COUNTIFS($A${M0}:$A${M1},A{r},$B${M0}:$B${M1},B{r})>1,"FEHLER: doppelt",'
                    f'IF(OR(C{r}="",D{r}=""),"FEHLER: Preis unvollständig",'
                    f'IF(AND(C{r}="€/Stange",E{r}=""),"FEHLER: Stangenlänge fehlt","OK"))))))')
add_dv(wpr, list_dv("lst_Profil", "Bitte Profil-ID aus Blatt 'Profile' wählen."), f"A{M0}:A{M1}")
add_dv(wpr, list_dv("lst_Werkstoff", "Bitte Werkstoff-ID aus Blatt 'Werkstoffe' wählen."), f"B{M0}:B{M1}")
add_dv(wpr, list_dv("lst_Preisbasis", "Bitte €/kg oder €/Stange wählen."), f"C{M0}:C{M1}")
add_dv(wpr, num_dv(0.01, 100000, msg="Preis in € (> 0)."), f"D{M0}:D{M1}")
add_dv(wpr, num_dv(500, 20000, whole=True, msg="Stangenlänge in mm (500–20000)."), f"E{M0}:E{M1}")
dd = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
dd.error = "Bitte gültiges Datum eingeben."; dd.showErrorMessage = True
add_dv(wpr, dd, f"F{M0}:F{M1}")
wpr.freeze_panes = "C5"

# ---------------------------------------------------------------- Maschinen
wm = wb.create_sheet("Maschinen", 5)
title(wm, "Maschinenparameter", "Endbeschnitt = Anschnitt am Stangenanfang; Greiferrest = nicht nutzbares Ende im Greifer; Schnittverlust = Materialverlust je Trennschnitt.")
cols_m = [("Maschinen-ID", 16, "in"), ("Bezeichnung", 26, "in"), ("Endbeschnitt [mm]", 12, "in"),
          ("Greiferrest [mm]", 12, "in"), ("Schnittverlust je Schnitt [mm]", 13, "in"),
          ("Max. Stangenlänge [mm]", 13, "in"), ("Stundensatz [€/h]", 12, "in"), ("Kontur klein bis [mm]", 11, "in"),
          ("Kontur mittel bis [mm]", 11, "in"), ("Abstand getrennte Schnitte [mm]", 12, "in"),
          ("Rüstzeit je Auftrag [min]", 11, "in"), ("Stangenwechsel je Stange [s]", 11, "in"), ("Bemerkung", 30, "in")]
header(wm, 4, cols_m)
style_rows(wm, M0, M1, cols_m, {3: "0", 4: "0", 5: "0.0", 6: "#,##0", 7: "#,##0.00", 10: "0.0", 11: "0", 12: "0"})
for i, row in enumerate([
    ("RL-1", "Rohrlaser 1 (Beispiel)", 20, 80, 0.3, 6500, 120.00, 10, 50, 2, 15, 30, "Beispielwerte"),
    ("RL-2", "Rohrlaser 2 (Beispiel)", 15, 120, 0.3, 8500, 140.00, 10, 50, 2, 20, 40, "Beispielwerte"),
], start=M0):
    for j, v in enumerate(row, start=1):
        wm.cell(row=i, column=j, value=v)
add_dv(wm, num_dv(0, 2000, msg="Wert in mm (0–2000)."), f"C{M0}:D{M1}")
add_dv(wm, num_dv(0, 20, msg="Schnittverlust in mm (0–20)."), f"E{M0}:E{M1}")
add_dv(wm, num_dv(500, 20000, whole=True, msg="Stangenlänge in mm (500–20000)."), f"F{M0}:F{M1}")
add_dv(wm, num_dv(0, 10000, msg="Stundensatz in €/h."), f"G{M0}:G{M1}")
add_dv(wm, num_dv(0.1, 5000, msg="Grenze in mm (größtes Hüllmaß der Kontur)."), f"H{M0}:I{M1}")
wm["H4"].comment = Comment("Innenkontur gilt als klein, wenn ihr größtes Hüllmaß ≤ diesem Wert ist.", "Kalkulation")
wm["I4"].comment = Comment("Bis zu diesem Hüllmaß mittel, darüber groß.", "Kalkulation")
add_dv(wm, num_dv(0, 100, msg="Abstand in mm (0–100)."), f"J{M0}:J{M1}")
add_dv(wm, num_dv(0, 1440, msg="Zeit in min."), f"K{M0}:K{M1}")
add_dv(wm, num_dv(0, 3600, msg="Zeit in s."), f"L{M0}:L{M1}")
wm["J4"].comment = Comment("Steg zwischen zwei getrennten Schnitten (wenn nicht beide Enden gerade sind). Je getrennter Trennung: 2 × Schnittverlust + dieser Abstand.", "Kalkulation")
wm["K4"].comment = Comment("Einmal je Auftrag, wird auf alle Teile umgelegt.", "Kalkulation")
wm["L4"].comment = Comment("Laden/Entladen je Stange, wird auf alle Teile umgelegt.", "Kalkulation")
dup = DataValidation(type="custom", formula1=f"=COUNTIF($A${M0}:$A${M1},A{M0})=1", allow_blank=True)
dup.error = "Diese Maschinen-ID existiert bereits."; dup.showErrorMessage = True
add_dv(wm, dup, f"A{M0}:A{M1}")
wm.freeze_panes = "B5"
wb.defined_names["lst_Maschine"] = DefinedName("lst_Maschine", attr_text=f"Maschinen!$A${M0}:$A${M1}")

# ---------------------------------------------------------------- Schneiddaten
wsd = wb.create_sheet("Schneiddaten", 6)
title(wsd, "Schneiddaten (Laserzeit)", "Je Maschine + Werkstoff gestaffelt nach Wandstärke: es gilt die Zeile mit der kleinsten 'Wandstärke bis', die ≥ der Profilwandstärke ist.")
cols_sd = [("Maschinen-ID", 16, "in"), ("Werkstoff-ID", 16, "in"), ("Wandstärke bis [mm]", 12, "in"),
           ("Vorschub groß / Trennschnitt [mm/min]", 13, "in"), ("Vorschub mittel [mm/min]", 12, "in"),
           ("Vorschub klein [mm/min]", 12, "in"), ("Einstechzeit [s]", 12, "in"), ("Vorschub Gravur [mm/min]", 12, "in"),
           ("Bemerkung", 30, "in")]
header(wsd, 4, cols_sd)
style_rows(wsd, M0, M1, cols_sd, {3: "0.0", 4: "#,##0", 5: "#,##0", 6: "#,##0", 7: "0.0", 8: "#,##0"})
sd_rows = []
for m in ("RL-1",):
    for w, data in (("S235-10219", [(2, 6000, 0.3), (3, 4500, 0.5), (5, 3000, 0.8)]),
                    ("S355-10219", [(2, 5800, 0.3), (3, 4300, 0.5), (5, 2800, 0.8)]),
                    ("1.4301", [(2, 5000, 0.4), (3, 3500, 0.6), (5, 2000, 1.0)]),
                    ("AW6060-T66", [(2, 7000, 0.3), (3, 5000, 0.5), (5, 3200, 0.8)])):
        for t, v, e in data:
            # Beispiel: mittlere Konturen 70 %, kleine 40 % des Vorschubs groß
            sd_rows.append((m, w, t, v, round(v * 0.7, -1), round(v * 0.4, -1), e, 10000, "Beispiel"))
for i, row in enumerate(sd_rows, start=M0):
    for j, v in enumerate(row, start=1):
        wsd.cell(row=i, column=j, value=v)
add_dv(wsd, list_dv("lst_Maschine", "Bitte Maschinen-ID aus Blatt 'Maschinen' wählen."), f"A{M0}:A{M1}")
add_dv(wsd, list_dv("lst_Werkstoff", "Bitte Werkstoff-ID aus Blatt 'Werkstoffe' wählen."), f"B{M0}:B{M1}")
add_dv(wsd, num_dv(0.1, 50, msg="Wandstärke in mm."), f"C{M0}:C{M1}")
add_dv(wsd, num_dv(1, 100000, msg="Vorschub in mm/min."), f"D{M0}:F{M1}")
add_dv(wsd, num_dv(0, 60, msg="Einstechzeit in s."), f"G{M0}:G{M1}")
add_dv(wsd, num_dv(1, 100000, msg="Vorschub in mm/min."), f"H{M0}:H{M1}")
wsd["H4"].comment = Comment("Leer = Vorschub groß wird verwendet.", "Kalkulation")
wsd["D4"].comment = Comment("Gilt für Anfangs-/Endschnitt und große Innenkonturen.", "Kalkulation")
wsd["E4"].comment = Comment("Leer = Vorschub groß wird verwendet.", "Kalkulation")
wsd["F4"].comment = Comment("Leer = Vorschub mittel bzw. groß wird verwendet.", "Kalkulation")
wsd.freeze_panes = "C5"

# ---------------------------------------------------------------- Gewinde
wg = wb.create_sheet("Gewinde", 7)
title(wg, "Gewinde / Kernlöcher", "ISO-Regelgewinde, Kernloch nach DIN 336. Der STEP-Import vergleicht Bohrungsdurchmesser mit dieser Tabelle (±0,05 mm).")
cols_g = [("Gewinde", 12, "in"), ("Kernloch Ø [mm]", 12, "in"), ("Zeit je Gewinde [s]", 12, "in"), ("Bemerkung", 36, "in")]
header(wg, 4, cols_g)
style_rows(wg, M0, M0 + 49, cols_g, {2: "0.0", 3: "0"})
for i, (g_, d_) in enumerate([("M3", 2.5), ("M4", 3.3), ("M5", 4.2), ("M6", 5.0), ("M8", 6.8), ("M10", 8.5),
                              ("M12", 10.2), ("M14", 12.0), ("M16", 14.0), ("M20", 17.5)], start=M0):
    wg.cell(row=i, column=1, value=g_)
    wg.cell(row=i, column=2, value=d_)
wg.cell(row=M0, column=4, value="Zeit je Gewinde für Kalkulation in Schritt 4")
add_dv(wg, num_dv(0.5, 100, msg="Kernloch in mm."), f"B{M0}:B{M0 + 49}")
add_dv(wg, num_dv(0, 3600, msg="Zeit in s."), f"C{M0}:C{M0 + 49}")
wg.freeze_panes = "B5"

# ---------------------------------------------------------------- Auftrag
wa = wb.create_sheet("Auftrag", 1)
title(wa, "Auftrag – Eingabe", "Kopf ausfüllen, dann Positionen ab Zeile 13. Status in Spalte W beachten. Beispielpositionen vor Gebrauch löschen (nur gelbe Zellen).")
kopf = [("Auftragsnummer", "A-2026-0001"), ("Kundennummer", "K-10045"), ("Kundenname", "Beispiel GmbH"),
        ("Datum", datetime.date(2026, 10, 15)), ("Maschine", "RL-1"), ("Standard-Stangenlänge [mm]", 6000)]
for i, (lab, val) in enumerate(kopf, start=4):
    wa.cell(row=i, column=1, value=lab).font = f_bold
    wa.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
    c = wa.cell(row=i, column=3, value=val)
    c.font = f_in; c.fill = fill_in; c.border = border
    wa.merge_cells(start_row=i, start_column=3, end_row=i, end_column=4)
wa["C7"].number_format = "DD.MM.YYYY"
wa["C9"].number_format = "#,##0"

mlook = lambda col: f'=IF($C$8="","",IFERROR(INDEX(Maschinen!${col}${M0}:${col}${M1},MATCH($C$8,lst_Maschine,0)),"?"))'
mk = [("Endbeschnitt [mm]", mlook("C"), "0"), ("Greiferrest [mm]", mlook("D"), "0"),
      ("Schnittverlust [mm]", mlook("E"), "0.0"), ("Max. Stangenlänge [mm]", mlook("F"), "#,##0"),
      ("Stundensatz [€/h]", mlook("G"), "#,##0.00"), ("Abstand getr. Schnitte [mm]", mlook("J"), "0.0")]
for i, (lab, fml, fmt) in enumerate(mk, start=4):
    wa.cell(row=i, column=6, value=lab).font = f_bold
    wa.merge_cells(start_row=i, start_column=6, end_row=i, end_column=7)
    c = wa.cell(row=i, column=8, value=fml)
    c.font = f_link; c.fill = fill_calc; c.border = border; c.number_format = fmt
wa["F3"] = "Maschinenwerte (aus Blatt 'Maschinen')"; wa["F3"].font = f_note
OR_ = f"{O0}:{{c}}{O1}"
rng = lambda c: f"${c}${O0}:${c}${O1}"
summ = [("Positionen", f'=COUNTIF({rng("W")},"?*")', "0"),
        ("davon FEHLER", f'=COUNTIF({rng("W")},"FEHLER*")', "0"),
        ("davon WARNUNG/HINWEIS", f'=COUNTIF({rng("W")},"WARNUNG*")+COUNTIF({rng("W")},"HINWEIS*")', "0"),
        ("Teile gesamt [Stk]", f'=SUM({rng("G")})', "#,##0"),
        ("Teilelänge gesamt [m]", f'=SUM({rng("P")})/1000', "#,##0.00"),
        ("Teilegewicht gesamt [kg]", f'=SUM({rng("X")})', "#,##0.00")]
for i, (lab, fml, fmt) in enumerate(summ, start=4):
    wa.cell(row=i, column=10, value=lab).font = f_bold
    wa.merge_cells(start_row=i, start_column=10, end_row=i, end_column=12)
    c = wa.cell(row=i, column=13, value=fml)
    c.font = f_calc; c.fill = fill_calc; c.border = border; c.number_format = fmt
wa["J3"] = "Übersicht"; wa["J3"].font = f_note

cols_a = [("Pos.", 6, "in"), ("Artikelnummer", 16, "in"), ("Bezeichnung", 22, "in"), ("Profil-ID", 18, "in"),
          ("Werkstoff-ID", 15, "in"), ("Teilelänge [mm]", 11, "in"), ("Stückzahl", 9, "in"),
          ("Schneidweg je Teil ohne Gravur [mm] (STEP)", 13, "in"), ("Einstiche je Teil (STEP)", 10, "in"),
          ("Stangenlänge abweichend [mm]", 12, "in"),
          ("Profilform", 13, "link"), ("Querschnitt [mm²]", 10, "link"), ("Wandstärke [mm]", 10, "link"),
          ("Metergewicht [kg/m]", 10, "calc"), ("Gewicht je Teil [kg]", 10, "calc"), ("Länge gesamt [mm]", 11, "calc"),
          ("Stangenlänge wirksam [mm]", 11, "calc"), ("Endbeschnitt wirksam [mm]", 10, "calc"),
          ("Greiferrest wirksam [mm]", 10, "calc"), ("Nutzlänge Stange [mm]", 11, "calc"),
          ("Preis hinterlegt", 9, "calc"), ("Vorschub groß [mm/min]", 10, "link"), ("Status", 34, "calc"), ("Gewicht Position [kg]", 11, "calc"),
          ("Anfangsschnitt gerade (STEP)", 10, "in"), ("Anfangsschnitt [mm] (STEP)", 11, "in"),
          ("Endschnitt gerade (STEP)", 10, "in"), ("Endschnitt [mm] (STEP)", 11, "in"),
          ("Innenkonturen [mm] (STEP)", 11, "in"), ("Gravur [mm] (STEP)", 10, "in"),
          ("Bohrungen (STEP)", 16, "in"), ("Gewinde-Vorschlag (STEP)", 18, "in"),
          ("Gewinde geprüft (manuell)", 16, "in"),
          ("Kontur klein [Stk] (STEP)", 9, "in"), ("Kontur klein [mm] (STEP)", 10, "in"),
          ("Kontur mittel [Stk] (STEP)", 9, "in"), ("Kontur mittel [mm] (STEP)", 10, "in"),
          ("Kontur groß [Stk] (STEP)", 9, "in"), ("Kontur groß [mm] (STEP)", 10, "in"),
          ("Eingeplant [Stk]", 9, "calc"), ("Nicht eingeplant [Stk]", 9, "calc")]
header(wa, 12, cols_a)
wa.row_dimensions[12].height = 60
style_rows(wa, O0, O1, cols_a, {6: "#,##0.0", 7: "#,##0", 8: "#,##0.0", 9: "0", 10: "#,##0", 12: "0.0",
                                 13: "0.00", 14: "0.000", 15: "0.000", 16: "#,##0", 17: "#,##0", 18: "0",
                                 19: "0", 20: "#,##0", 22: "#,##0", 24: "#,##0.00", 26: "#,##0.0", 28: "#,##0.0", 29: "#,##0.0", 30: "#,##0.0",
                                 34: "0", 35: "#,##0.0", 36: "0", 37: "#,##0.0", 38: "0", 39: "#,##0.0"})
P = lambda col: f"Profile!${col}${M0}:${col}${M1}"
DET = lambda col: f"Stangenbelegung!${col}${D0}:${col}${D1}"
SD = lambda col: f"Schneiddaten!${col}${M0}:${col}${M1}"
for r in range(O0, O1 + 1):
    mp = f"MATCH($D{r},lst_Profil,0)"
    mw = f"MATCH($E{r},lst_Werkstoff,0)"
    wa[f"K{r}"] = f'=IF($D{r}="","",IFERROR(INDEX({P("B")},{mp}),""))'
    wa[f"L{r}"] = f'=IF($D{r}="","",IFERROR(INDEX({P("K")},{mp})+0,""))'
    wa[f"M{r}"] = f'=IF($D{r}="","",IFERROR(INDEX({P("E")},{mp})+0,""))'
    wa[f"N{r}"] = f'=IF(OR(L{r}="",$E{r}=""),"",IFERROR(L{r}*INDEX(Werkstoffe!$E${M0}:$E${M1},{mw})/1000,""))'
    wa[f"O{r}"] = f'=IF(OR(N{r}="",$F{r}=""),"",N{r}*$F{r}/1000)'
    wa[f"P{r}"] = f'=IF(OR($F{r}="",$G{r}=""),"",$F{r}*$G{r})'
    wa[f"Q{r}"] = f'=IF($D{r}="","",IF($J{r}<>"",$J{r},$C$9))'
    wa[f"R{r}"] = f'=IF($D{r}="","",IFERROR(IF(INDEX({P("G")},{mp})<>"",INDEX({P("G")},{mp}),$H$4),$H$4))'
    wa[f"S{r}"] = f'=IF($D{r}="","",IFERROR(IF(INDEX({P("H")},{mp})<>"",INDEX({P("H")},{mp}),$H$5),$H$5))'
    wa[f"T{r}"] = f'=IF(OR(Q{r}="",NOT(ISNUMBER(R{r})),NOT(ISNUMBER(S{r}))),"",Q{r}-R{r}-S{r})'
    wa[f"U{r}"] = (f'=IF(OR($D{r}="",$E{r}=""),"",IF(COUNTIFS(Preise!$A${M0}:$A${M1},$D{r},'
                   f'Preise!$B${M0}:$B${M1},$E{r},Preise!$I${M0}:$I${M1},"OK")>0,"ja","nein"))')
    thr = (f'_xlfn.MINIFS({SD("C")},{SD("A")},$C$8,{SD("B")},$E{r},{SD("C")},">="&M{r})')
    wa[f"V{r}"] = (f'=IF(OR(M{r}="",$E{r}="",$C$8=""),"",IF({thr}=0,"",'
                   f'SUMIFS({SD("D")},{SD("A")},$C$8,{SD("B")},$E{r},{SD("C")},{thr})))')
    wa[f"X{r}"] = f'=IF(OR(O{r}="",$G{r}=""),"",O{r}*$G{r})'
    wa[f"AN{r}"] = f'=IF($D{r}="","",COUNTIFS({DET("E")},ROW()))'
    wa[f"AO{r}"] = f'=IF(OR($G{r}="",COUNTA({DET("A")})=0),"",$G{r}-AN{r})'
    wa[f"W{r}"] = (
        f'=IF(AND($D{r}="",$E{r}="",$F{r}="",$G{r}=""),"",'
        f'IF(ISNA({mp}),"FEHLER: Profil unbekannt",'
        f'IF(ISNA({mw}),"FEHLER: Werkstoff unbekannt",'
        f'IF(NOT(ISNUMBER(L{r})),"FEHLER: Querschnitt fehlt (Profil)",'
        f'IF(OR($F{r}="",$F{r}<=0),"FEHLER: Teilelänge fehlt",'
        f'IF(OR($G{r}="",$G{r}<1),"FEHLER: Stückzahl fehlt",'
        f'IF(NOT(ISNUMBER($H$4)),"FEHLER: Maschine fehlt/unbekannt",'
        f'IF(Q{r}>$H$7,"FEHLER: Stange länger als Maschine zulässt",'
        f'IF($F{r}+$H$6>T{r},"FEHLER: Teil länger als Nutzlänge",'
        f'IF(N(AO{r})>0,"WARNUNG: nicht vollständig eingeplant",'
        f'IF(U{r}="nein","WARNUNG: kein gültiger Materialpreis",'
        f'IF($H{r}="","HINWEIS: Schneidweg fehlt",'
        f'IF(V{r}="","HINWEIS: keine Schneiddaten",'
        f'IF(AND($AF{r}<>"",$AG{r}=""),"HINWEIS: Gewinde-Vorschlag prüfen","OK"))))))))))))))')

ex = [(10, "4711-001", "Holm links", "RR 42,4x2,0", "S235-10219", 1250, 24, 410.5, 4, None),
      (20, "4711-002", "Querstrebe", "RR 42,4x2,0", "S235-10219", 860, 40, 302.0, 2, None),
      (30, "4711-003", "Rahmen oben", "QR 40x40x3", "S235-10219", 2150, 12, 520.0, 3, None),
      (40, "4711-004", "Stütze", "RE 60x40x3", "S355-10219", 1795.5, 16, 618.2, 4, None),
      (50, "4711-005", "Anschlag", "RR 48,3x2,6", "S235-10219", 345, 60, 255.0, 2, None)]
for i, row in enumerate(ex, start=O0):
    for j, v in enumerate(row, start=1):
        if v is not None:
            wa.cell(row=i, column=j, value=v)

add_dv(wa, list_dv("lst_Maschine", "Bitte Maschinen-ID aus Blatt 'Maschinen' wählen."), "C8")
add_dv(wa, num_dv(500, 20000, whole=True, msg="Stangenlänge in mm (500–20000)."), "C9")
dd = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
dd.error = "Bitte gültiges Datum eingeben."; dd.showErrorMessage = True
add_dv(wa, dd, "C7")
add_dv(wa, list_dv("lst_Profil", "Bitte Profil-ID aus Blatt 'Profile' wählen."), f"D{O0}:D{O1}")
add_dv(wa, list_dv("lst_Werkstoff", "Bitte Werkstoff-ID aus Blatt 'Werkstoffe' wählen."), f"E{O0}:E{O1}")
add_dv(wa, num_dv(1, 20000, msg="Teilelänge in mm (1–20000)."), f"F{O0}:F{O1}")
add_dv(wa, num_dv(1, 100000, whole=True, msg="Stückzahl als ganze Zahl ≥ 1."), f"G{O0}:G{O1}")
add_dv(wa, num_dv(0, 1000000, msg="Schneidweg in mm (≥ 0)."), f"H{O0}:H{O1}")
add_dv(wa, num_dv(0, 1000, whole=True, msg="Anzahl Einstiche als ganze Zahl ≥ 0."), f"I{O0}:I{O1}")
add_dv(wa, num_dv(500, 20000, whole=True, msg="Stangenlänge in mm (500–20000)."), f"J{O0}:J{O1}")
add_dv(wa, num_dv(1, 99999, whole=True, msg="Positionsnummer als ganze Zahl."), f"A{O0}:A{O1}")
wa["H12"].comment = Comment("Wird beim STEP-Import automatisch gefüllt, sonst manuell eintragen.", "Kalkulation")
wa["J12"].comment = Comment("Leer = Standard-Stangenlänge aus C9.", "Kalkulation")
add_dv(wa, list_dv("lst_JaNein", "Bitte ja oder nein wählen."), f"Y{O0}:Y{O1}")
add_dv(wa, list_dv("lst_JaNein", "Bitte ja oder nein wählen."), f"AA{O0}:AA{O1}")
add_dv(wa, num_dv(0, 100000, msg="Schnittlänge in mm (≥ 0)."), f"Z{O0}:Z{O1}")
add_dv(wa, num_dv(0, 100000, msg="Schnittlänge in mm (≥ 0)."), f"AB{O0}:AB{O1}")
wa["Y12"].comment = Comment("Gerade = Schnitt senkrecht zur Rohrachse. Zwei benachbarte gerade Enden teilen sich in der Optimierung einen Schnitt.", "Kalkulation")

wrng = f"W{O0}:W{O1}"
for prefix, color in (("FEHLER", "F8CBAD"), ("WARNUNG", "FFD966"), ("HINWEIS", "FFF2CC"), ("OK", "C6EFCE")):
    wa.conditional_formatting.add(wrng, FormulaRule(formula=[f'LEFT($W{O0},{len(prefix)})="{prefix}"'],
                                                    fill=PatternFill("solid", fgColor=color)))
wa.freeze_panes = "E13"
add_dv(wa, num_dv(0, 1000000, msg="Länge in mm (≥ 0)."), f"AC{O0}:AD{O1}")
gv = DataValidation(type="custom", allow_blank=True,
                    formula1=f'=ISNUMBER(SEARCH("xM",AG{O0}))')
gv.error = "Format: Anzahl x Gewinde, mehrere mit Semikolon, z. B. 4xM6; 2xM8"
gv.errorTitle = "Ungültige Eingabe"; gv.showErrorMessage = True
add_dv(wa, gv, f"AG{O0}:AG{O1}")
wa["AF12"].comment = Comment("Nur Vorschlag aus dem Kernlochdurchmesser (Blatt 'Gewinde'). Wird nicht gerechnet.", "Kalkulation")
wa["AG12"].comment = Comment("Nach manueller Prüfung eintragen, z. B. 4xM6; 2xM8. Nur dieser Wert wird gerechnet.", "Kalkulation")
wa.conditional_formatting.add(f"AF{O0}:AF{O1}", FormulaRule(formula=[f'AND($AF{O0}<>"",$AG{O0}="")'],
                                                            fill=PatternFill("solid", fgColor="FFD966")))
add_dv(wa, num_dv(0, 10000, whole=True, msg="Anzahl als ganze Zahl ≥ 0."), f"AH{O0}:AH{O1}")
add_dv(wa, num_dv(0, 10000, whole=True, msg="Anzahl als ganze Zahl ≥ 0."), f"AJ{O0}:AJ{O1}")
add_dv(wa, num_dv(0, 10000, whole=True, msg="Anzahl als ganze Zahl ≥ 0."), f"AL{O0}:AL{O1}")
for c_ in ("AI", "AK", "AM"):
    add_dv(wa, num_dv(0, 1000000, msg="Länge in mm (≥ 0)."), f"{c_}{O0}:{c_}{O1}")
wa["AH12"].comment = Comment("Innenkonturen nach größtem Hüllmaß; Grenzen je Maschine im Blatt 'Maschinen' (Spalten H/I).", "Kalkulation")
wa.auto_filter.ref = f"A12:AO{O1}"

# ---------------------------------------------------------------- Schnittplan (Stangen)
wsp = wb.create_sheet("Schnittplan", 2)
title(wsp, "Schnittplan je Stange",
      "Wird von der Optimierung gefüllt (Schaltfläche 'Optimieren' im Blatt 'Auftrag'). Graue Spalten rechnen selbst und prüfen jede Stange gegen Überbelegung.")
cols_sp = [("Stange Nr.", 8, "in"), ("Profil-ID", 18, "in"), ("Werkstoff-ID", 15, "in"), ("Stangenlänge [mm]", 10, "in"),
           ("Endbeschnitt [mm]", 9, "in"), ("Greiferrest [mm]", 9, "in"), ("Teile [Stk]", 8, "calc"),
           ("Summe Teilelängen [mm]", 11, "calc"), ("Schnittverluste + Abstände [mm]", 11, "in"),
           ("Belegt bis [mm]", 10, "calc"), ("Rest / Verschnitt inkl. Greiferrest [mm]", 11, "calc"),
           ("Ausnutzung", 9, "calc"), ("Prüfung", 13, "calc"), ("Trennungen [Stk]", 9, "in"),
           ("davon gemeinsamer Schnitt", 10, "in"), ("Belegung (Pos. × Länge)", 60, "in"),
           ("Stangengewicht [kg]", 10, "calc"), ("Materialkosten Stange [€]", 11, "calc")]
header(wsp, 4, cols_sp)
style_rows(wsp, S0, S1, cols_sp, {4: "#,##0", 5: "0", 6: "0", 7: "0", 8: "#,##0.0", 9: "#,##0.0", 10: "#,##0.0",
                                   11: "#,##0.0", 12: "0.0%", 14: "0", 15: "0", 17: "#,##0.00", 18: "#,##0.00"})
for r in range(S0, S1 + 1):
    g = f'IF($A{r}="","",'
    wsp[f"G{r}"] = f'={g}COUNTIFS({DET("A")},$A{r}))'
    wsp[f"H{r}"] = f'={g}SUMIFS({DET("F")},{DET("A")},$A{r}))'
    wsp[f"J{r}"] = f'={g}$E{r}+H{r}+$I{r})'
    wsp[f"K{r}"] = f'={g}$D{r}-J{r})'
    wsp[f"L{r}"] = f'={g}H{r}/$D{r})'
    wsp[f"M{r}"] = (f'={g}IF(OR(J{r}+$F{r}>$D{r}+0.001,_xlfn.MAXIFS({DET("H")},{DET("A")},$A{r})>$D{r}-$F{r}+0.001),'
                    f'"ÜBERBELEGT","OK"))')
    kgm = (f'INDEX(Profile!$K${M0}:$K${M1},MATCH($B{r},lst_Profil,0))'
           f'*INDEX(Werkstoffe!$E${M0}:$E${M1},MATCH($C{r},lst_Werkstoff,0))/1000')
    wsp[f"Q{r}"] = f'={g}IFERROR($D{r}/1000*{kgm},""))'
    pk = lambda col: f'Preise!${col}${M0}:${col}${M1}'
    crit = f'{pk("A")},$B{r},{pk("B")},$C{r},{pk("I")},"OK"'
    wsp[f"R{r}"] = (f'={g}IF(COUNTIFS({crit},{pk("C")},"€/kg")>0,N(Q{r})*SUMIFS({pk("D")},{crit}),'
                    f'IF(COUNTIFS({crit},{pk("C")},"€/Stange")>0,'
                    f'SUMIFS({pk("D")},{crit})*$D{r}/SUMIFS({pk("E")},{crit}),0)))')
wsp.conditional_formatting.add(f"M{S0}:M{S1}", FormulaRule(formula=[f'$M{S0}="ÜBERBELEGT"'], fill=PatternFill("solid", fgColor="F8CBAD")))
wsp.conditional_formatting.add(f"M{S0}:M{S1}", FormulaRule(formula=[f'$M{S0}="OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))
wsp.freeze_panes = "B5"
wsp.page_setup.orientation = "landscape"
wsp.page_setup.fitToWidth = 1
wsp.page_setup.fitToHeight = 0
wsp.sheet_properties.pageSetUpPr.fitToPage = True
wsp.print_title_rows = "4:4"

# ---------------------------------------------------------------- Stangenbelegung (Teile je Stange)
wdt = wb.create_sheet("Stangenbelegung", 3)
title(wdt, "Stangenbelegung je Teil",
      "Ein Eintrag je Teil, in Reihenfolge auf der Stange. Von/Bis ab Stangenanfang. Gemeinsamer Schnitt = beide Enden gerade.")
cols_dt = [("Stange Nr.", 8, "in"), ("Platz", 6, "in"), ("Pos.", 6, "in"), ("Artikelnummer", 22, "in"),
           ("Zeile Auftrag", 8, "in"), ("Teilelänge [mm]", 10, "in"), ("Von [mm]", 10, "in"), ("Bis [mm]", 10, "in"),
           ("Lage", 9, "in"), ("Trennung zum Vorgänger", 16, "in"), ("Eingesparter Schneidweg [mm]", 11, "in"),
           ("Eingesparte Einstiche", 9, "in")]
header(wdt, 4, cols_dt)
for c_ in range(1, len(cols_dt) + 1):
    wdt.cell(row=4, column=c_).border = border
wdt.freeze_panes = "A5"
wdt.print_title_rows = "4:4"

# ---------------------------------------------------------------- Auswertung
wau = wb.create_sheet("Auswertung", 4)
title(wau, "Auswertung Materialoptimierung",
      "Je Profil/Werkstoff/Stangenlänge: optimierte Einteilung gegen einfache manuelle Einteilung (Reihenfolge der Positionen, Stange voll, dann nächste).")
wau.column_dimensions["A"].width = 20

# ---------------------------------------------------------------- Kalkulation
wk = wb.create_sheet("Kalkulation", 2)
title(wk, "Kalkulation je Teil und Auftrag",
      "Material aus dem Schnittplan (Verschnitt anteilig nach Länge umgelegt), Laserzeit aus Schneidweg je Konturklasse, Gewinde nur aus 'Gewinde geprüft'.")
for i, (lab, val, fmt) in enumerate([("Zuschlag Material", 0.10, "0.0%"), ("Zuschlag Fertigung", 0.15, "0.0%"),
                                      ("Gewinn", 0.10, "0.0%")], start=4):
    wk.cell(row=i, column=1, value=lab).font = f_bold
    wk.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
    c = wk.cell(row=i, column=3, value=val)
    c.font = f_in; c.fill = fill_in; c.border = border; c.number_format = fmt
wk["A3"] = "Zuschläge (Beispielwerte)"; wk["A3"].font = f_note
add_dv(wk, num_dv(0, 5, msg="Zuschlag als Anteil, z. B. 10 % = 0,10."), "C4:C6")
mk_k = lambda col: f'IFERROR(INDEX(Maschinen!${col}${M0}:${col}${M1},MATCH(Auftrag!$C$8,lst_Maschine,0))+0,0)'
info = [("Rüstzeit je Auftrag [min]", f"={mk_k('K')}", "0"),
        ("Stangenwechsel je Stange [s]", f"={mk_k('L')}", "0"),
        ("Stangen laut Schnittplan", f"=COUNT(Schnittplan!$A${S0}:$A${S1})", "0"),
        ("Rüsten + Stangenwechsel [€]", "=(F4*60+F5*F6)/3600*N(Auftrag!$H$8)", "#,##0.00"),
        ("Teile gesamt [Stk]", f"=SUM($C${O0}:$C${O1})", "#,##0")]
for i, (lab, fml, fmt) in enumerate(info, start=4):
    wk.cell(row=i, column=4, value=lab).font = f_bold
    wk.merge_cells(start_row=i, start_column=4, end_row=i, end_column=5)
    c = wk.cell(row=i, column=6, value=fml)
    c.font = f_calc; c.fill = fill_calc; c.border = border; c.number_format = fmt
wk["D3"] = "Auftragskosten (umgelegt auf alle Teile)"; wk["D3"].font = f_note
tot = [("Material [€]", f"=SUM($X${O0}:$X${O1})"),
       ("Laser + Gewinde [€]", f"=SUM($Y${O0}:$Y${O1})"),
       ("Rüsten + Stangenwechsel [€]", "=F7"),
       ("Herstellkosten [€]", f"=SUM($V${O0}:$V${O1})"),
       ("Verkaufspreis Auftrag [€]", f"=SUM($U${O0}:$U${O1})")]
for i, (lab, fml) in enumerate(tot, start=4):
    wk.cell(row=i, column=8, value=lab).font = f_bold
    wk.merge_cells(start_row=i, start_column=8, end_row=i, end_column=10)
    c = wk.cell(row=i, column=11, value=fml)
    c.font = Font(name=F, bold=True); c.fill = fill_calc; c.border = border; c.number_format = "#,##0.00"
wk["H3"] = "Summen Auftrag"; wk["H3"].font = f_note

cols_k = [("Pos.", 6, "link"), ("Artikelnummer", 18, "link"), ("Stückzahl", 8, "link"),
          ("Materialkosten Gruppe [€]", 11, "calc"), ("Längenanteil an Gruppe", 9, "calc"),
          ("Material je Teil [€]", 10, "calc"), ("Wandstärke-Stufe [mm]", 9, "calc"),
          ("Vorschub groß [mm/min]", 9, "link"), ("Vorschub mittel [mm/min]", 9, "link"),
          ("Vorschub klein [mm/min]", 9, "link"), ("Vorschub Gravur [mm/min]", 9, "link"), ("Einstechzeit [s]", 8, "link"),
          ("Laserzeit brutto je Teil [s]", 10, "calc"), ("Einsparung gemeinsame Schnitte je Teil [s]", 11, "calc"),
          ("Laserzeit netto je Teil [s]", 10, "calc"), ("Gewindezeit je Teil [s]", 10, "calc"),
          ("Laser + Gewinde je Teil [€]", 10, "calc"), ("Rüsten + Stangenwechsel je Teil [€]", 11, "calc"),
          ("Herstellkosten je Teil [€]", 11, "calc"), ("Verkaufspreis je Teil [€]", 11, "calc"),
          ("Verkaufspreis Position [€]", 12, "calc"), ("Herstellkosten Position [€]", 12, "calc"), ("Hinweis", 30, "calc"),
          ("Material Position [€]", 11, "calc"), ("Laser + Gewinde Position [€]", 11, "calc")]
header(wk, 12, cols_k)
wk.row_dimensions[12].height = 60
style_rows(wk, O0, O1, cols_k, {3: "#,##0", 4: "#,##0.00", 5: "0.0%", 6: "#,##0.000", 7: "0.0", 8: "#,##0", 9: "#,##0",
                                 10: "#,##0", 11: "#,##0", 12: "0.0", 13: "#,##0.0", 14: "#,##0.0", 15: "#,##0.0",
                                 16: "#,##0.0", 17: "#,##0.000", 18: "#,##0.000", 19: "#,##0.000", 20: "#,##0.000",
                                 21: "#,##0.00", 22: "#,##0.00", 24: "#,##0.00", 25: "#,##0.00"})
for r in range(O0, O1 + 1):
    a = f"Auftrag!$%s${r}"
    g = f'IF($A{r}="","",'
    wk[f"A{r}"] = f'=IF(Auftrag!$D{r}="","",Auftrag!$A{r})'
    wk[f"B{r}"] = f'={g}Auftrag!$B{r})'
    wk[f"C{r}"] = f'={g}N(Auftrag!$G{r}))'
    wk[f"D{r}"] = (f'={g}SUMIFS(Schnittplan!$R${S0}:$R${S1},Schnittplan!$B${S0}:$B${S1},Auftrag!$D{r},'
                   f'Schnittplan!$C${S0}:$C${S1},Auftrag!$E{r},Schnittplan!$D${S0}:$D${S1},Auftrag!$Q{r}))')
    wk[f"E{r}"] = (f'={g}IFERROR(N(Auftrag!$P{r})/SUMIFS(Auftrag!$P${O0}:$P${O1},Auftrag!$D${O0}:$D${O1},Auftrag!$D{r},'
                   f'Auftrag!$E${O0}:$E${O1},Auftrag!$E{r},Auftrag!$Q${O0}:$Q${O1},Auftrag!$Q{r}),0))')
    wk[f"F{r}"] = f'={g}IF($C{r}=0,0,$D{r}*$E{r}/$C{r}))'
    crit = f'{SD("A")},Auftrag!$C$8,{SD("B")},Auftrag!$E{r}'
    wk[f"G{r}"] = f'={g}_xlfn.MINIFS({SD("C")},{crit},{SD("C")},">="&N(Auftrag!$M{r})))'
    sv = lambda col: f'SUMIFS({SD(col)},{crit},{SD("C")},$G{r})'
    wk[f"H{r}"] = f'={g}IF(N($G{r})=0,"",{sv("D")}))'
    wk[f"I{r}"] = f'={g}IF($H{r}="","",IF({sv("E")}>0,{sv("E")},$H{r})))'
    wk[f"J{r}"] = f'={g}IF($H{r}="","",IF({sv("F")}>0,{sv("F")},$I{r})))'
    wk[f"K{r}"] = f'={g}IF($H{r}="","",IF({sv("H")}>0,{sv("H")},$H{r})))'
    wk[f"L{r}"] = f'={g}IF($H{r}="","",{sv("G")}))'
    hm = lambda col: f"N(Auftrag!${col}{r})"
    wk[f"M{r}"] = (f'={g}IF($H{r}="","",60*(({hm("H")}-{hm("AI")}-{hm("AK")})/$H{r}+{hm("AI")}/$J{r}+{hm("AK")}/$I{r}'
                   f'+{hm("AD")}/$K{r})+{hm("I")}*$L{r}))')
    wk[f"N{r}"] = (f'={g}IF(OR($H{r}="",$C{r}=0),0,(60*SUMIFS({DET("K")},{DET("E")},ROW())/$H{r}'
                   f'+SUMIFS({DET("L")},{DET("E")},ROW())*$L{r})/$C{r}))')
    wk[f"O{r}"] = f'={g}IF($M{r}="","",MAX(0,$M{r}-$N{r})))'
    T = f'";"&SUBSTITUTE(UPPER(Auftrag!$AG{r})," ","")&";"'
    K = f'"X"&UPPER(Gewinde!$A${M0}:$A${M0 + 49})&";"'
    wk[f"P{r}"] = (f'={g}IFERROR(SUMPRODUCT(--("0"&TRIM(RIGHT(SUBSTITUTE(LEFT({T}&{K},SEARCH({K},{T}&{K})-1),";",'
                   f'REPT(" ",50)),50))),Gewinde!$C${M0}:$C${M0 + 49}),0))')
    wk[f"Q{r}"] = f'={g}(N($O{r})+$P{r})/3600*N(Auftrag!$H$8))'
    wk[f"R{r}"] = f'={g}IF($F$8=0,0,$F$7/$F$8))'
    wk[f"S{r}"] = f'={g}$F{r}*(1+$C$4)+($Q{r}+$R{r})*(1+$C$5))'
    wk[f"T{r}"] = f'={g}$S{r}*(1+$C$6))'
    wk[f"U{r}"] = f'={g}$T{r}*$C{r})'
    wk[f"V{r}"] = f'={g}$S{r}*$C{r})'
    wk[f"X{r}"] = f'={g}$F{r}*$C{r})'
    wk[f"Y{r}"] = f'={g}$Q{r}*$C{r})'
    wk[f"W{r}"] = (f'={g}IF(N(Auftrag!$AO{r})>0,"nicht vollständig eingeplant",IF($D{r}=0,"kein Schnittplan oder Preis",'
                   f'IF($H{r}="","keine Schneiddaten",IF(AND(Auftrag!$AF{r}<>"",Auftrag!$AG{r}=""),"Gewinde ungeprüft","")))))')
wk.freeze_panes = "C13"

# ---------------------------------------------------------------- STEP-Import (Protokoll)
# Spalten wie LOG_KOPF in tools/step_import.py (geprüft in tests/test_step_import.py)
LOG_KOPF = ["Zeitpunkt", "Datei", "Körper", "Profilform", "Abm. A [mm]", "Abm. B [mm]", "Wandstärke [mm]",
            "Querschnitt [mm²]", "Länge [mm]", "Schneidweg ohne Gravur [mm]", "Einstiche",
            "Anfangsschnitt gerade", "Anfangsschnitt [mm]", "Endschnitt gerade", "Endschnitt [mm]",
            "Innenkonturen [mm]", "Innenkonturen [Stk]", "Gravur [mm]", "Gravur [Stk]", "Bohrungen",
            "Gewinde-Vorschlag", "Kontur klein [Stk]", "Kontur klein [mm]", "Kontur mittel [Stk]",
            "Kontur mittel [mm]", "Kontur groß [Stk]", "Kontur groß [mm]", "Profil-ID", "Zeile Auftrag", "Hinweise"]
wlog = wb.create_sheet("STEP-Import")
wlog["A1"] = "STEP-Import Protokoll"
wlog["A1"].font = Font(name=F, size=14, bold=True)
wlog["A2"] = "Messwerte je Teil. Schneidweg = Konturlänge auf der Außenhaut, beide Endschnitte voll gezählt."
wlog["A2"].font = Font(name=F, italic=True, color="595959", size=9)
for i, k in enumerate(LOG_KOPF, start=1):
    c = wlog.cell(4, i, k)
    c.font = f_head
    c.fill = fill_head
    wlog.column_dimensions[c.column_letter].width = 14
wlog.column_dimensions["B"].width = 34
wlog.column_dimensions["AD"].width = 50
wlog.freeze_panes = "A5"

for s in wb.worksheets:
    s.sheet_view.zoomScale = 90

# ---------------------------------------------------------------- VBA-Projekt und Schaltflächen
import glob, tempfile
from vba_projekt import vba_projekt, xlsm_packen, schaltflaechen_xml
wb.code_name = "DieseArbeitsmappe"
for i, s in enumerate(wb.worksheets, start=1):
    s.sheet_properties.codeName = f"Tabelle{i}"
module = {}
for datei in sorted(glob.glob(os.path.join(os.path.dirname(__file__), "..", "vba", "*.bas"))):
    code = open(datei, encoding="utf-8").read()
    name = os.path.splitext(os.path.basename(datei))[0]
    module[name] = "\n".join(z for z in code.splitlines() if not z.startswith("Attribute VB_Name"))
dokumente = [("DieseArbeitsmappe", "", True)] + [(f"Tabelle{i}", "", False) for i in range(1, len(wb.worksheets) + 1)]
knoepfe = [("Optimieren", "Optimieren", (13, 3, 16, 5), "1F4E78"),
           ("StepImportieren", "STEP importieren", (13, 6, 16, 8), "548235")]
with tempfile.TemporaryDirectory() as tmp:
    roh = os.path.join(tmp, "roh.xlsx")
    wb.save(roh)
    xlsm_packen(roh, OUT, vba_projekt(module, dokumente), {"Auftrag": schaltflaechen_xml(knoepfe)})
print("saved", OUT)
