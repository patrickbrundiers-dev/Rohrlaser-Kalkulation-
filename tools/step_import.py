"""STEP-Import in die Kalkulationsmappe.

Liest STEP-Dateien, misst je Teil Profil, Länge, Schneidweg und Konturen und
trägt die Teile als neue Positionen in das Blatt 'Auftrag' ein. Alle Messwerte
werden zusätzlich im Blatt 'STEP-Import' protokolliert.

Aufruf:
    python step_import.py MAPPE.xlsx DATEI_ODER_ORDNER [...] [--werkstoff ID] [--stueck N]
                          [--neue-profile] [--erneut]

Vor dem Schreiben wird eine Sicherung MAPPE.bak.xlsx angelegt.
"""
import argparse
import datetime
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent))
from step_geometrie import (KERNLOCH, KLEIN_BIS, MITTEL_BIS, analysiere_datei, bohrungen_text,  # noqa: E402
                            gewinde_vorschlag, konturklassen)

M0, M1 = 5, 304           # Stammdatenzeilen
O0, O1 = 13, 512          # Auftragspositionen
TOL_MASS = 0.15           # mm, Toleranz beim Profilabgleich
LOG = "STEP-Import"
LOG_KOPF = ["Zeitpunkt", "Datei", "Körper", "Profilform", "Abm. A [mm]", "Abm. B [mm]", "Wandstärke [mm]",
            "Querschnitt [mm²]", "Länge [mm]", "Schneidweg ohne Gravur [mm]", "Einstiche",
            "Anfangsschnitt gerade", "Anfangsschnitt [mm]", "Endschnitt gerade", "Endschnitt [mm]",
            "Innenkonturen [mm]", "Innenkonturen [Stk]", "Gravur [mm]", "Gravur [Stk]", "Bohrungen",
            "Gewinde-Vorschlag", "Kontur klein [Stk]", "Kontur klein [mm]", "Kontur mittel [Stk]",
            "Kontur mittel [mm]", "Kontur groß [Stk]", "Kontur groß [mm]", "Profil-ID", "Zeile Auftrag", "Hinweise"]


def lies_konturgrenzen(wb):
    """Grenzen klein/mittel der im Auftragskopf gewählten Maschine (Blatt 'Maschinen', Spalten H/I)."""
    maschine = wb["Auftrag"]["C8"].value
    ws = wb["Maschinen"]
    for r in range(M0, M1 + 1):
        if maschine and ws.cell(r, 1).value == maschine:
            k, m = ws.cell(r, 8).value, ws.cell(r, 9).value
            if isinstance(k, (int, float)) and isinstance(m, (int, float)) and 0 < k < m:
                return float(k), float(m), ""
            break
    return KLEIN_BIS, MITTEL_BIS, f"Konturgrenzen Vorgabe {KLEIN_BIS:g}/{MITTEL_BIS:g} mm (Maschine ohne Werte)"


def lies_gewinde(wb):
    if "Gewinde" not in wb.sheetnames:
        return KERNLOCH
    ws = wb["Gewinde"]
    tab = [(ws.cell(r, 1).value, ws.cell(r, 2).value) for r in range(M0, M1 + 1)]
    tab = [(g, float(d)) for g, d in tab if g and isinstance(d, (int, float))]
    return tab or KERNLOCH


def lies_profile(ws):
    profile = []
    for r in range(M0, M1 + 1):
        pid = ws.cell(r, 1).value
        if pid in (None, ""):
            continue
        profile.append(dict(id=pid, form=ws.cell(r, 2).value, a=ws.cell(r, 3).value or 0,
                            b=ws.cell(r, 4).value or 0, t=ws.cell(r, 5).value or 0))
    return profile


def finde_profil(teil, profile):
    for p in profile:
        if p["form"] != teil.profilform or abs(p["t"] - teil.wandstaerke) > TOL_MASS:
            continue
        if teil.profilform in ("Rundrohr", "Quadratrohr"):
            if abs(p["a"] - teil.abm_a) <= TOL_MASS:
                return p["id"]
        elif teil.profilform == "Rechteckrohr":
            pa, pb = max(p["a"], p["b"]), min(p["a"], p["b"])
            if abs(pa - teil.abm_a) <= TOL_MASS and abs(pb - teil.abm_b) <= TOL_MASS:
                return p["id"]
    return None


def profil_anlegen(ws, teil):
    for r in range(M0, M1 + 1):
        if ws.cell(r, 1).value in (None, ""):
            ws.cell(r, 1, teil.profil_id_vorschlag)
            ws.cell(r, 2, teil.profilform)
            ws.cell(r, 3, round(teil.abm_a, 2))
            if teil.profilform == "Rechteckrohr":
                ws.cell(r, 4, round(teil.abm_b, 2))
            ws.cell(r, 5, round(teil.wandstaerke, 2))
            ws.cell(r, 6, round(teil.flaeche, 1))
            ws.cell(r, 9, "aus STEP-Import angelegt, bitte prüfen")
            return r
    raise RuntimeError("Blatt 'Profile' ist voll")


def log_blatt(wb):
    if LOG in wb.sheetnames:
        return wb[LOG]
    ws = wb.create_sheet(LOG)
    ws["A1"] = "STEP-Import Protokoll"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = "Messwerte je Teil. Schneidweg = Konturlänge auf der Außenhaut, beide Endschnitte voll gezählt."
    ws["A2"].font = Font(name="Arial", italic=True, color="595959", size=9)
    for i, k in enumerate(LOG_KOPF, start=1):
        c = ws.cell(4, i, k)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        ws.column_dimensions[c.column_letter].width = 14
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["X"].width = 50
    ws.freeze_panes = "A5"
    return ws


def step_dateien(pfade):
    out = []
    for p in map(Path, pfade):
        if p.is_dir():
            out += sorted(x for x in p.iterdir() if x.suffix.lower() in (".step", ".stp"))
        else:
            out.append(p)
    return out


def importiere(mappe, pfade, werkstoff=None, stueck=None, neue_profile=False, erneut=False, sicherung=True):
    mappe = Path(mappe)
    if stueck is not None and (int(stueck) != stueck or stueck < 1):
        raise ValueError("Stückzahl muss eine ganze Zahl ≥ 1 sein")
    if sicherung:
        shutil.copy2(mappe, mappe.with_suffix(".bak" + mappe.suffix))
    wb = load_workbook(mappe, keep_vba=mappe.suffix.lower() == ".xlsm")
    wa, wp = wb["Auftrag"], wb["Profile"]
    wl = log_blatt(wb)
    profile = lies_profile(wp)
    gewinde = lies_gewinde(wb)
    klein_bis, mittel_bis, grenz_hinweis = lies_konturgrenzen(wb)

    vorhandene = {wa.cell(r, 2).value for r in range(O0, O1 + 1) if wa.cell(r, 2).value}
    pos = [wa.cell(r, 1).value for r in range(O0, O1 + 1) if isinstance(wa.cell(r, 1).value, (int, float))]
    naechste_pos = int(max(pos, default=0)) + 10
    belegt = [r for r in range(O0, O1 + 1) if any(wa.cell(r, c).value not in (None, "") for c in range(1, 11))]
    zeile = (max(belegt) + 1) if belegt else O0
    jetzt = datetime.datetime.now().replace(microsecond=0)
    ergebnis = []

    for datei in step_dateien(pfade):
        try:
            teile = analysiere_datei(datei)
        except Exception as e:  # Datei defekt o. ä.: protokollieren, weitermachen
            wl.append([jetzt, datei.name] + [""] * (len(LOG_KOPF) - 3) + [f"FEHLER: {e}"])
            ergebnis.append((datei.name, None, f"FEHLER: {e}"))
            continue
        for t in teile:
            artikel = t.datei if len(teile) == 1 else f"{t.datei}-K{t.koerper}"
            hinweise = list(t.hinweise)
            vorschlag = gewinde_vorschlag(t.bohrungen, gewinde)
            klassen = konturklassen(t.innen_liste, klein_bis, mittel_bis)
            if grenz_hinweis and t.innen_liste:
                hinweise.append(grenz_hinweis)
            ziel = ""
            pid = finde_profil(t, profile) if t.profilform else None
            if t.profilform and not pid:
                if neue_profile and t.profil_id_vorschlag:
                    profil_anlegen(wp, t)
                    profile = lies_profile(wp)
                    pid = t.profil_id_vorschlag
                    hinweise.append("Profil neu in Stammdaten angelegt")
                else:
                    pid = t.profil_id_vorschlag
                    hinweise.append("Profil fehlt in Stammdaten")
            if artikel in vorhandene and not erneut:
                hinweise.append("übersprungen: Artikel steht schon im Auftrag")
            elif t.laenge <= 0 or not t.profilform:
                hinweise.append("nicht übernommen: kein Rohrteil erkannt")
            elif zeile > O1:
                hinweise.append("nicht übernommen: Auftragstabelle voll")
            else:
                wa.cell(zeile, 1, naechste_pos)
                wa.cell(zeile, 2, artikel)
                wa.cell(zeile, 4, pid)
                if werkstoff:
                    wa.cell(zeile, 5, werkstoff)
                wa.cell(zeile, 6, round(t.laenge, 1))
                if stueck is not None:
                    wa.cell(zeile, 7, int(stueck))
                wa.cell(zeile, 8, round(t.schneidweg, 1))
                wa.cell(zeile, 9, t.konturen)
                if t.ende_a:
                    wa.cell(zeile, 25, "ja" if t.ende_a_gerade else "nein")
                    wa.cell(zeile, 26, round(t.ende_a, 1))
                if t.ende_b:
                    wa.cell(zeile, 27, "ja" if t.ende_b_gerade else "nein")
                    wa.cell(zeile, 28, round(t.ende_b, 1))
                wa.cell(zeile, 29, round(t.innen, 1))
                wa.cell(zeile, 30, round(t.gravur, 1))
                if t.bohrungen:
                    wa.cell(zeile, 31, bohrungen_text(t.bohrungen))
                if vorschlag:
                    wa.cell(zeile, 32, vorschlag)
                for i, k in enumerate(("klein", "mittel", "gross")):
                    wa.cell(zeile, 34 + 2 * i, klassen[k][0])
                    wa.cell(zeile, 35 + 2 * i, round(klassen[k][1], 1))
                ziel = zeile
                vorhandene.add(artikel)
                naechste_pos += 10
                zeile += 1
            wl.append([jetzt, datei.name, t.koerper, t.profilform, round(t.abm_a, 2), round(t.abm_b, 2) or None,
                       round(t.wandstaerke, 2), round(t.flaeche, 1), round(t.laenge, 1), round(t.schneidweg, 1),
                       t.konturen, "ja" if t.ende_a_gerade else "nein", round(t.ende_a, 1),
                       "ja" if t.ende_b_gerade else "nein", round(t.ende_b, 1), round(t.innen, 1),
                       t.innen_konturen, round(t.gravur, 1), t.gravur_konturen, bohrungen_text(t.bohrungen),
                       vorschlag, klassen["klein"][0], round(klassen["klein"][1], 1),
                       klassen["mittel"][0], round(klassen["mittel"][1], 1),
                       klassen["gross"][0], round(klassen["gross"][1], 1), pid, ziel, "; ".join(hinweise)])
            ergebnis.append((artikel, ziel or None, "; ".join(hinweise) or "OK"))

    for row in wl.iter_rows(min_row=5):
        for c in row:
            c.font = Font(name="Arial")
        row[0].number_format = "DD.MM.YYYY HH:MM"
    wb.calculation.fullCalcOnLoad = True
    wb.save(mappe)
    return ergebnis


def main():
    ap = argparse.ArgumentParser(description="STEP-Dateien in die Kalkulationsmappe übernehmen")
    ap.add_argument("mappe")
    ap.add_argument("pfade", nargs="+", help="STEP-Dateien oder Ordner")
    ap.add_argument("--werkstoff", help="Werkstoff-ID für alle importierten Teile")
    ap.add_argument("--stueck", type=int, help="Stückzahl für alle importierten Teile")
    ap.add_argument("--neue-profile", action="store_true", help="fehlende Profile in 'Profile' anlegen")
    ap.add_argument("--erneut", action="store_true", help="Artikel auch dann eintragen, wenn schon vorhanden")
    a = ap.parse_args()
    if a.stueck is not None and a.stueck < 1:
        ap.error("--stueck muss eine ganze Zahl ≥ 1 sein")
    for artikel, zeile, info in importiere(a.mappe, a.pfade, a.werkstoff, a.stueck, a.neue_profile, a.erneut):
        print(f"{artikel:40s} Zeile {zeile or '-':>4}  {info}")


if __name__ == "__main__":
    main()
