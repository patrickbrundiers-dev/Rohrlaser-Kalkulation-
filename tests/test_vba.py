"""Prüft die VBA-Makros der Mappe gegen die Python-Fassungen.

1. Optimierung (tools/optimierung.py): Je Testauftrag wird die Mappe mit Python und mit
   dem VBA-Makro in LibreOffice gerechnet. Schnittplan, Stangenbelegung und Auswertung
   müssen übereinstimmen, und die Formelprüfung im Schnittplan (Spalte M) muss für jede
   Stange OK zeigen.
2. STEP-Import (tools/step_import.py): Die Messwerte der Testteile (tools/step_messen.py,
   wie STEP-Import.exe) werden per VBA übernommen. Auftrag, Profile und Protokoll müssen
   mit dem Python-Import übereinstimmen.

3. Bedienung und Ausgabe: Merker "Optimierung veraltet", Werkstattpapier und Angebot füllen,
   neuer Auftrag, Stammdaten aus einer anderen Mappe übernehmen.

4. Staffelpreise (tools/staffeln.py): Jede Mengenstufe wird vom VBA-Makro als eigener Auftrag
   gerechnet. Stangenzahl je Stufe gegen die Python-Optimierung, Preise gegen die Kalkulation,
   Wiederherstellung des Auftrags.

Aufruf: python tests/test_vba.py [optimierung|import|bedienung|staffeln]   (benötigt LibreOffice, für 2. auch OCP)
"""
import random
import shutil
import sys
import tempfile
from pathlib import Path

HIER = Path(__file__).parent
sys.path.insert(0, str(HIER.parent / "tools"))
sys.path.insert(0, str(HIER))
from openpyxl import load_workbook  # noqa: E402
from optimierung import optimiere_mappe  # noqa: E402
import staffeln  # noqa: E402
from lo_makro import LibreOffice  # noqa: E402

MAPPE = HIER.parent / "mappe" / "Rohrlaser_Kalkulation_v02.xlsm"
PROFILE = ["RR 42,4x2,0", "RR 48,3x2,6", "QR 40x40x3", "RE 60x40x3", "RE 120x60x4"]
WERKSTOFFE = ["S235-10219", "S355-10219", "1.4301"]


def norm(v):
    if v == "":
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    return v


def gleich(a, b):
    a, b = norm(a), norm(b)
    if isinstance(a, float) and isinstance(b, float):
        return abs(a - b) <= 0.1001   # Python und VBA runden Grenzfälle (x,x5) verschieden
    return a == b


def lies(pfad, werte=False):
    wb = load_workbook(pfad, data_only=werte)
    wsp, wdt, wau = wb["Schnittplan"], wb["Stangenbelegung"], wb["Auswertung"]
    sp = [[wsp.cell(r, c).value for c in (1, 2, 3, 4, 5, 6, 9, 14, 15, 16)] for r in range(5, 1005)
          if wsp.cell(r, 1).value is not None]
    dt = [[wdt.cell(r, c).value for c in range(1, 13)] for r in range(5, 10005) if wdt.cell(r, 1).value is not None]
    au = [[wau.cell(r, c).value for c in range(1, 18)] for r in range(4, wau.max_row + 1)]
    while au and all(v in (None, "") for v in au[-1]):
        au.pop()
    pruef = [wsp.cell(r, 13).value for r in range(5, 5 + len(sp))] if werte else None
    return sp, dt, au, pruef


def vergleiche(name, py, vb):
    for teil, a, b in zip(("Schnittplan", "Stangenbelegung", "Auswertung"), py[:3], vb[:3]):
        assert len(a) == len(b), f"{name} {teil}: {len(a)} Zeilen (Python) gegen {len(b)} (VBA)"
        for i, (za, zb) in enumerate(zip(a, b)):
            for j, (x, y) in enumerate(zip(za, zb)):
                if isinstance(x, str) and isinstance(y, str) and x.startswith("=") and y.startswith("="):
                    x, y = x.replace("_xlfn.", ""), y.replace("_xlfn.", "")
                assert gleich(x, y), f"{name} {teil} Zeile {i + 1} Spalte {j + 1}: Python {x!r}, VBA {y!r}"


def zufallsauftrag(wb, rnd):
    wa = wb["Auftrag"]
    for r in range(13, 513):
        for c in list(range(1, 11)) + list(range(25, 29)):
            wa.cell(r, c).value = None
    wa["C9"] = rnd.choice([3000, 6000, 6500])
    for i in range(rnd.randint(1, 8)):
        r = 13 + i
        stange = wa["C9"].value
        L = round(rnd.uniform(40, stange - 101), 1) if rnd.random() < 0.15 else round(rnd.uniform(40, 1900), 1)
        werte = {1: (i + 1) * 10 if rnd.random() < 0.9 else None, 2: f"T-{i}" if rnd.random() < 0.9 else None,
                 4: rnd.choice(PROFILE) if rnd.random() < 0.95 else "XX 1x1", 5: rnd.choice(WERKSTOFFE[:2]),
                 6: L, 7: rnd.randint(1, 25) if rnd.random() < 0.95 else None}
        if rnd.random() < 0.3:
            werte[8] = round(rnd.uniform(100, 900), 1)
        if rnd.random() < 0.1:
            werte[10] = 6500
        for c, v in werte.items():
            wa.cell(r, c, v)
        if rnd.random() < 0.25:
            wa.cell(r, 25, "nein")
        if rnd.random() < 0.25:
            wa.cell(r, 27, "nein")
        if rnd.random() < 0.4:
            wa.cell(r, 26, round(rnd.uniform(50, 200), 1))
            wa.cell(r, 28, round(rnd.uniform(50, 200), 1))


def sonderfaelle(wb, rnd):
    wa = wb["Auftrag"]
    zusatz = [(60, "Z-zu-lang", "zu lang", "RR 42,4x2,0", "S235-10219", 5950, 1, None, None, None),
              (70, "Z-profil", "Profil fehlt", "ZZ 99x9", "S235-10219", 500, 2, None, None, None),
              (80, "Z-schraeg", "Gehrung", "RR 42,4x2,0", "S235-10219", 700, 7, 300, 2, None),
              (90, "Z-8m", "lange Stange", "QR 40x40x3", "S235-10219", 2600, 5, 400, 2, 8000),
              (None, None, None, "QR 40x40x3", "S235-10219", 777, 3, None, None, None),
              (110, "Z-stk", None, "QR 40x40x3", "S235-10219", 777, 2.5, None, None, None),
              (120, "Z-auto", "Profil aus Namen", "RR 99x9", "S235-10219", 500, 2, None, None, None)]
    for i, werte in enumerate(zusatz):
        for c, v in enumerate(werte, start=1):
            wa.cell(18 + i, c, v)
    wa.cell(20, 25, "nein"); wa.cell(20, 26, 160.0); wa.cell(20, 27, "ja"); wa.cell(20, 28, 133.2)
    wb["Maschinen"]["F5"] = 8500


def schwer(wb, rnd):
    """Eine Gruppe mit 10 bis 25 Einzelteilen: hier ist First/Best-Fit oft eine Stange schlechter
    als die exakte Suche."""
    wa = wb["Auftrag"]
    for r in range(13, 513):
        for c in list(range(1, 11)) + list(range(25, 29)):
            wa.cell(r, c).value = None
    for i in range(rnd.randint(10, 25)):
        for c, v in enumerate([(i + 1) * 10, f"S-{i}", None, "QR 40x40x3", "S235-10219",
                               float(round(rnd.uniform(300, 2500))), 1], start=1):
            wa.cell(13 + i, c, v)
        if rnd.random() < 0.3:
            wa.cell(13 + i, 25, "nein")
        if rnd.random() < 0.3:
            wa.cell(13 + i, 27, "nein")


def test_optimierung(lo, d):
    rnd = random.Random(2610)
    faelle = [("Beispielauftrag", None), ("Sonderfälle", sonderfaelle)]
    faelle += [(f"Zufall {i + 1}", zufallsauftrag) for i in range(25)]
    faelle += [(f"Exakt {i + 1}", schwer) for i in range(12)]
    if True:
        for name, aendern in faelle:
            ein = d / "ein.xlsm"
            shutil.copy(MAPPE, ein)
            if aendern:
                wb = load_workbook(ein, keep_vba=True)
                aendern(wb, rnd)
                wb.save(ein)
            py = d / "py.xlsm"
            shutil.copy(ein, py)
            optimiere_mappe(py, sicherung=False)
            aus = d / "vba.xlsx"
            text = lo.ausfuehren(ein, "Optimierung.OptimiereTest", aus)
            assert not text.startswith("FEHLER"), f"{name}: {text}"
            vb = lies(aus)
            vergleiche(name, lies(py), vb)
            pruef = lies(aus, werte=True)[3]
            assert all(p == "OK" for p in pruef), f"{name}: Prüfung Schnittplan {set(pruef)}"
            bew = {z[16] for z in vb[2][1:] if z[16]}
            print(f"OK     {name}: {len(vb[0])} Stangen, {len(vb[1])} Teile {sorted(bew)}")
    print("OK: VBA-Optimierung = Python-Optimierung")


# ------------------------------------------------------------------ STEP-Import
STEP = HIER / "step"
AUFTRAG_SPALTEN = list(range(1, 11)) + list(range(25, 40))


def lies_import(pfad):
    wb = load_workbook(pfad)
    wa, wp, wl = wb["Auftrag"], wb["Profile"], wb["STEP-Import"]
    au = [[wa.cell(r, c).value for c in AUFTRAG_SPALTEN] for r in range(13, 513)]
    pr = [[wp.cell(r, c).value for c in range(1, 10)] for r in range(5, 305)]
    lg = [[wl.cell(r, c).value for c in range(2, 31)] for r in range(5, wl.max_row + 1)]
    while lg and all(v in (None, "") for v in lg[-1]):
        lg.pop()
    return au, pr, lg


def test_import(lo, d):
    from step_import import importiere   # benötigt OCP
    from step_messen import main as messen
    mess = d / "messwerte.txt"
    umlaut = d / "Halter_äöü_Ø.step"   # Dateiname mit Umlauten (Messwerte sind UTF-8)
    shutil.copy(STEP / "T01_RR42-4x2_L1250.step", umlaut)
    dateien = [str(STEP / "T07_QR40x3_L1000_Gewinde_Gravur.step"), str(STEP), str(umlaut), str(d / "fehlt.step")]
    assert messen(["--ausgabe", str(mess)] + dateien) == 1     # fehlt.step -> Fehlerzeile
    # (Name, Mappe ändern, [(Werkstoff, Stückzahl, neue Profile, erneut), ...] nacheinander)
    faelle = [("Import neu", None, [("S235-10219", 10, True, False)]),
              ("Import ohne Profile anlegen", None, [("", 0, False, False)]),
              ("Import zweimal", None, [("S235-10219", 5, False, False), ("", 0, False, False)]),
              ("Import erneut", None, [("", 3, True, False), ("1.4301", 0, True, True)]),
              ("Maschine ohne Konturgrenzen", "grenzen", [("", 0, True, False)])]
    for name, aendern, laeufe in faelle:
        ein = d / "ein.xlsm"
        shutil.copy(MAPPE, ein)
        if aendern == "grenzen":
            wb = load_workbook(ein, keep_vba=True)
            wb["Maschinen"]["H5"] = None
            wb.save(ein)
        py = d / "py.xlsm"
        shutil.copy(ein, py)
        for wst, stk, neu, erneut in laeufe:
            importiere(py, dateien, wst or None, stk or None, neu, erneut, sicherung=False)
        aus = d / "vba.xlsx"
        text = lo.ausfuehren(ein, [("StepImport.ImportTest", (str(mess), wst, stk, neu, erneut))
                                   for wst, stk, neu, erneut in laeufe], aus)
        assert not text.startswith("FEHLER"), f"{name}: {text}"
        for teil, a, b in zip(("Auftrag", "Profile", "Protokoll"), lies_import(py), lies_import(aus)):
            assert len(a) == len(b), f"{name} {teil}: {len(a)} Zeilen (Python) gegen {len(b)} (VBA)"
            for i, (za, zb) in enumerate(zip(a, b)):
                for j, (x, y) in enumerate(zip(za, zb)):
                    assert gleich(x, y), f"{name} {teil} Zeile {i + 1} Feld {j + 1}: Python {x!r}, VBA {y!r}"
        print(f"OK     {name}")
    print("OK: VBA-STEP-Import = Python-STEP-Import")


# ------------------------------------------------------------------ Bedienung und Ausgabe
def test_bedienung(lo, d):
    ein = d / "ein.xlsm"
    shutil.copy(MAPPE, ein)
    aus = d / "vba.xlsx"
    lo.ausfuehren(ein, [("Optimierung.OptimiereTest", ()), ("Bedienung.GeaendertTest", ()),
                        ("Optimierung.OptimiereTest", ()), ("Ausgabe.AusgabeTest", ())], aus)
    r = lo.rueckgaben
    assert not r[0].startswith("FEHLER"), r[0]
    assert r[1] == "nein,ja", f"Merker Optimierung veraltet: {r[1]}"
    assert r[3].startswith("OK"), r[3]
    wb = load_workbook(aus, data_only=True)
    assert wb["Listen"]["E5"].value == "nein"
    assert wb["Auftrag"]["C10"].value == "aktuell", wb["Auftrag"]["C10"].value
    # Werkstattpapier: Materialliste und Blöcke decken alle Stangen und Teile ab
    wsp, wdt, wwp = wb["Schnittplan"], wb["Stangenbelegung"], wb["Werkstattpapier"]
    stangen = [wsp.cell(r_, 1).value for r_ in range(5, 1005) if wsp.cell(r_, 1).value is not None]
    teile = sum(1 for r_ in range(5, 10005) if wdt.cell(r_, 1).value is not None)
    zeilen = [[wwp.cell(r_, c).value for c in range(1, 11)] for r_ in range(1, wwp.max_row + 1)]
    mat = []
    i = next(k for k, z in enumerate(zeilen) if z[0] == "Profil-ID") + 1
    while zeilen[i][0]:
        mat.append(zeilen[i]); i += 1
    assert sum(z[5] for z in mat) == len(stangen), (mat, len(stangen))
    bloecke = [z[0] for z in zeilen if isinstance(z[0], str) and z[0].startswith("Stange")]
    n_stangen, n_teile, faktor = 0, 0, 0
    for z in zeilen:
        if isinstance(z[0], str) and z[0].startswith("Stange"):
            faktor = int(z[0].split("(")[1].split()[0]) if "(" in z[0] else 1
            n_stangen += faktor
        elif isinstance(z[0], (int, float)) and z[4] is not None:
            n_teile += faktor
    assert n_stangen == len(stangen) and n_teile == teile, (n_stangen, len(stangen), n_teile, teile)
    print(f"OK     Werkstattpapier: {len(bloecke)} Blöcke für {len(stangen)} Stangen, {teile} Teile")
    # Stangenbild: gleiche Muster wie im Werkstattpapier, je Teil ein Abschnitt mit Schnittlinie links
    wsb = wb["Stangenbild"]
    assert int(r[3].split()[-1]) > 0, r[3]
    n_stangen, n_teile, n_bloecke = 0, 0, 0
    for zz in range(3, wsb.max_row + 1):
        kopf_ = wsb.cell(zz, 2).value
        if not (isinstance(kopf_, str) and kopf_.startswith("Stange")):
            continue
        n_bloecke += 1
        faktor = int(kopf_.split("(")[1].split()[0]) if "(" in kopf_.split("|")[0] else 1
        n_stangen += faktor
        anzahl = int(kopf_.split("|")[4].split()[0])
        abschnitte = sum(1 for c in range(2, 122) if (wsb.cell(zz + 1, c).border.left.style or "") in ("medium", "thick"))
        assert abschnitte == anzahl, (kopf_, abschnitte)
        n_teile += faktor * anzahl
    assert n_bloecke == len(bloecke) and n_stangen == len(stangen) and n_teile == teile, (n_bloecke, n_stangen, n_teile)
    print(f"OK     Stangenbild: {n_bloecke} Bilder für {n_stangen} Stangen, {n_teile} Teile")
    # Angebot: eine Zeile je Einzelteil, je Satz ein Block mit Satzpreis; Preise = Kalkulation (auf Cent)
    wan, wk, wa = wb["Angebot"], wb["Kalkulation"], wb["Auftrag"]
    pos = [r_ for r_ in range(13, 513) if wa.cell(r_, 4).value not in (None, "")]
    ep = {r_: round(wk.cell(r_, 20).value, 2) for r_ in pos}
    je = {r_: (wa.cell(r_, 44).value or 1) for r_ in pos}
    satz = {r_: (wa.cell(r_, 43).value or "") for r_ in pos}
    zeilen, gesehen = [], set()
    for r_ in pos:
        if not satz[r_]:
            zeilen.append(("einzel", r_))
        elif satz[r_] not in gesehen:
            gesehen.add(satz[r_])
            zeilen.append(("satz", satz[r_]))
            zeilen += [("teil", k_) for k_ in pos if satz[k_] == satz[r_]]
    assert gesehen, "Beispielauftrag ohne Satz – Satzblock wird nicht geprüft"
    summe = 0.0
    bezeichnung = ("Rundrohr", "Quadratrohr", "Rechteckrohr")
    for k, (art, ref) in enumerate(zeilen):
        z = [wan.cell(14 + k, c).value for c in range(1, 7)]
        if art == "einzel":
            assert z[1] == wa.cell(ref, 2).value and abs(z[4] - ep[ref]) < 1e-9, (z, ep[ref])
            assert abs(z[5] - round(ep[ref] * wa.cell(ref, 7).value, 2)) < 1e-9
            assert z[2].startswith(bezeichnung) and ", L = " in z[2] and "inkl. Material" in z[2], z[2]
            summe += z[5]
        elif art == "satz":
            glieder = [k_ for k_ in pos if satz[k_] == ref]
            preis = round(sum(ep[k_] * je[k_] for k_ in glieder), 2)
            saetze = max(round(wa.cell(k_, 7).value / je[k_]) for k_ in glieder)
            assert z[2] == f"Satz {ref}" and z[3] == saetze, (z, saetze)
            assert abs(z[4] - preis) < 1e-9 and abs(z[5] - round(preis * saetze, 2)) < 1e-9, (z, preis)
            summe += z[5]
        else:
            assert z[1] == wa.cell(ref, 2).value, z
            assert z[4] in (None, "") and z[5] in (None, ""), f"Satzteil mit Preis: {z}"
            assert z[2].strip().startswith(f"{int(je[ref])} x"), z[2]
    zs = 14 + len(zeilen) + 1
    assert wan.cell(zs, 4).value == "Summe netto [€]" and abs(wan.cell(zs, 6).value - summe) < 0.005
    assert abs(wan.cell(zs + 2, 6).value - (summe + round(summe * 0.19, 2))) < 0.005
    print(f"OK     Angebot: {len(zeilen)} Zeilen für {len(pos)} Positionen und {len(gesehen)} Satz/Sätze, "
          f"{summe:.2f} € netto, Text: {wan.cell(14, 3).value}")

    # Materialpreis je Position im Auftrag (Spalte AP) schlägt Preise und Tonnenpreis
    preis_ein = d / "preis.xlsm"
    shutil.copy(MAPPE, preis_ein)
    wbp = load_workbook(preis_ein, keep_vba=True)
    wbp["Auftrag"]["AP13"] = 9.5     # Zeile 13: RR 42,4x2,0 / S235-10219 / 6000 mm
    wbp.save(preis_ein)
    text = lo.ausfuehren(preis_ein, "Optimierung.OptimiereTest", d / "preis.xlsx")
    assert not text.startswith("FEHLER"), text
    wbp = load_workbook(d / "preis.xlsx", data_only=True)
    wsp2, wa2 = wbp["Schnittplan"], wbp["Auftrag"]
    stangen = [(wsp2.cell(r_, 2).value, wsp2.cell(r_, 4).value, wsp2.cell(r_, 18).value)
               for r_ in range(5, 1005) if wsp2.cell(r_, 1).value is not None]
    mit = [k for k in stangen if k[0] == "RR 42,4x2,0" and k[1] == 6000]
    ohne = [k for k in stangen if k[0] != "RR 42,4x2,0"]
    assert mit and all(abs(k[2] - 9.5 * k[1] / 1000) < 0.005 for k in mit), mit[:3]
    assert ohne and all(abs(k[2] - 9.5 * k[1] / 1000) > 0.005 for k in ohne), ohne[:3]
    assert wa2["U13"].value == "ja"
    print(f"OK     Materialpreis im Auftrag: {len(mit)} Stangen zu {mit[0][2]:.2f} €")

    # Neuer Auftrag
    text = lo.ausfuehren(ein, [("Optimierung.OptimiereTest", ()), ("Ausgabe.AusgabeTest", ()),
                               ("Bedienung.NeuerAuftragTest", ())], d / "neu.xlsx")
    assert text == "OK A-2026-0002", text
    wb = load_workbook(d / "neu.xlsx", data_only=True)
    wa = wb["Auftrag"]
    assert all(wa.cell(r_, c).value in (None, "") for r_ in range(13, 513) for c in list(range(1, 11)) + list(range(25, 40)) + [42, 43, 44])
    assert wa["C8"].value == "RL-1" and wa["C9"].value == 6000 and wa["C6"].value in (None, "")
    assert wb["Schnittplan"]["A5"].value in (None, "") and wb["Stangenbelegung"]["A5"].value in (None, "")
    assert wa["C10"].value == "noch nicht optimiert", wa["C10"].value
    assert wb["Angebot"]["A14"].value in (None, "") and wb["Werkstattpapier"]["A5"].value in (None, "")
    assert wb["Staffeln"]["A5"].value in (None, "") and wb["Staffelpreise"]["A5"].value in (None, "")
    print("OK     Neuer Auftrag")

    # Stammdaten aus einer anderen Mappe übernehmen
    quelle = d / "quelle.xlsm"
    shutil.copy(MAPPE, quelle)
    wb = load_workbook(quelle, keep_vba=True)
    wb["Profile"]["A10"] = "QR 50x50x4"; wb["Profile"]["B10"] = "Quadratrohr"
    wb["Profile"]["C10"] = 50; wb["Profile"]["E10"] = 4
    wb["Werkstoffe"]["E7"] = 7.93
    wb["Angebot"]["A1"] = "Laser GmbH"
    wb["Kalkulation"]["C6"] = 0.2
    wb.save(quelle)
    text = lo.ausfuehren(ein, ("Bedienung.StammdatenTest"), d / "stamm.xlsx", (str(quelle),))
    assert text.startswith("Stammdaten übernommen"), text
    wb = load_workbook(d / "stamm.xlsx")
    wp = wb["Profile"]
    assert [wp.cell(10, c).value for c in (1, 2, 3, 5)] == ["QR 50x50x4", "Quadratrohr", 50, 4]
    assert str(wp["K10"].value).startswith("="), "Formelspalte überschrieben"
    assert wb["Werkstoffe"]["E7"].value == 7.93 and wb["Werkstoffe"]["A7"].value == "1.4301"
    assert wb["Angebot"]["A1"].value == "Laser GmbH" and wb["Kalkulation"]["C6"].value == 0.2
    assert load_workbook(d / "stamm.xlsx", data_only=True)["Listen"]["E5"].value == "ja"
    print("OK     Stammdaten übernehmen")


def test_staffeln(lo, d):
    """Staffelpreise: je Mengenstufe ein eigener Rechenlauf, Auftrag danach wiederhergestellt."""
    ein = d / "staffel.xlsm"
    shutil.copy(MAPPE, ein)
    aus = d / "staffel.xlsx"
    lo.ausfuehren(ein, [("Optimierung.OptimiereTest", ()), ("Staffel.StaffelTest", ()),
                        ("Ausgabe.AusgabeTest", ())], aus)
    r = lo.rueckgaben
    assert not r[0].startswith("FEHLER"), r[0]
    assert "gerechnet" in r[1], r[1]
    assert r[2].startswith("OK"), r[2]

    vorlage = load_workbook(MAPPE)
    erwartet = staffeln.szenarien(vorlage)
    assert erwartet, "Beispielmappe ohne Staffeln"
    wb = load_workbook(aus, data_only=True)
    wf, wa = wb["Staffelpreise"], wb["Auftrag"]

    # Auftrag unverändert: Stückzahlen wie in der Vorlage, Optimierung wieder aktuell
    for zeile in range(13, 513):
        assert wa.cell(zeile, 7).value == vorlage["Auftrag"].cell(zeile, 7).value, f"Stückzahl Zeile {zeile}"
    assert wb["Listen"]["E5"].value == "nein" and wa["C10"].value == "aktuell", wa["C10"].value

    zm, zf, zg = (wb["Kalkulation"][c].value for c in ("C4", "C5", "C6"))
    zeilen = [[wf.cell(z, c).value for c in range(1, 15)] for z in range(5, 205) if wf.cell(z, 1).value is not None]
    assert len(zeilen) == len(erwartet), f"{len(zeilen)} Zeilen gegen {len(erwartet)} Staffeln"
    for z, (schl, menge, ist_satz, mengen) in zip(zeilen, erwartet):
        assert staffeln._txt(z[0]) == schl and z[2] == ("Satz" if ist_satz else "Stück") and z[3] == menge, z
        assert z[4] == sum(mengen.values()), (z, mengen)
        # Stangenzahl gegen die Python-Optimierung des gleichen Szenarios
        pfad = d / f"sz_{schl}_{int(menge)}.xlsx"
        wbs = load_workbook(MAPPE)
        staffeln.setze(wbs, mengen)
        wbs.save(pfad)
        optimiere_mappe(pfad, sicherung=False)
        wbs = load_workbook(pfad)
        py = sum(1 for zz in range(5, 1005) if wbs["Schnittplan"].cell(zz, 1).value is not None)
        assert z[5] == py, f"Staffel {schl}/{menge:.0f}: VBA {z[5]} Stangen, Python {py}"
        # Herstellkosten = Material x (1 + Zuschlag Material) + (Laser + Rüsten) x (1 + Zuschlag Fertigung)
        assert abs(z[9] - (z[6] * (1 + zm) + (z[7] + z[8]) * (1 + zf))) < 0.05, z
        assert abs(z[10] - z[9] / menge) < 0.005, z
        assert abs(z[11] - round(z[9] * (1 + zg) / menge, 2)) < 0.011, z
        assert abs(z[12] - round(z[11] * menge, 2)) < 0.005, z
    # Stückpreis sinkt mit der Menge (Rüsten wird auf mehr Einheiten verteilt)
    je_schluessel = {}
    for z in zeilen:
        je_schluessel.setdefault(staffeln._txt(z[0]), []).append((z[3], z[10]))
    for schl, stufen in je_schluessel.items():
        stufen.sort()
        for (m1, p1), (m2, p2) in zip(stufen, stufen[1:]):
            assert p2 <= p1 + 1e-9, f"Staffel {schl}: {m2:.0f} Stück teurer je Einheit als {m1:.0f}"
    # Preistreppe im Angebot
    wan = wb["Angebot"]
    kopf = next(zz for zz in range(14, wan.max_row + 1) if wan.cell(zz, 1).value == "Staffelpreise")
    for i, z in enumerate(zeilen):
        a = [wan.cell(kopf + 2 + i, c).value for c in range(1, 7)]
        if z[2] == "Satz":
            assert a[2] == f"Satz {z[0]} ({z[1]})", (a, z)
        else:
            assert staffeln._txt(a[0]) == staffeln._txt(z[0]) and a[2] == z[1], (a, z)
        assert a[3] == z[3] and abs(a[4] - z[11]) < 1e-9 and abs(a[5] - z[12]) < 1e-9, (a, z)
    print(f"OK     Staffelpreise: {len(zeilen)} Stufen, Preise je Einheit "
          + ", ".join(f"{z[3]:.0f}x{z[2]} {z[11]:.2f} €" for z in zeilen))


def main():
    welche = sys.argv[1:] or ["optimierung", "import", "bedienung", "staffeln"]
    with tempfile.TemporaryDirectory() as d, LibreOffice() as lo:
        if "optimierung" in welche:
            test_optimierung(lo, Path(d))
        if "import" in welche:
            test_import(lo, Path(d))
        if "bedienung" in welche:
            test_bedienung(lo, Path(d))
        if "staffeln" in welche:
            test_staffeln(lo, Path(d))


if __name__ == "__main__":
    main()
