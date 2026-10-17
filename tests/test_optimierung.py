"""Prüft die Stangenoptimierung: Schnittregel, Drehen, Überbelegung, nicht einplanbare Teile, Mappe.

Aufruf: python tests/test_optimierung.py
"""
import random
import shutil
import sys
import tempfile
from pathlib import Path

HIER = Path(__file__).parent
sys.path.insert(0, str(HIER.parent / "tools"))
from openpyxl import load_workbook  # noqa: E402
from optimierung import (Gruppe, Teil, manuell, optimiere, optimiere_mappe, pruefe, untergrenze,  # noqa: E402
                         verbrauch, gemeinsame_schnitte)

MAPPE = HIER.parent / "mappe" / "Rohrlaser_Kalkulation_v02.xlsm"


def gruppe(teile, stange=1000):
    # Endbeschnitt 20, Greiferrest 80 -> belegbar bis 920; Schnittverlust 0,3; Abstand 2
    g = Gruppe("P", "W", stange, 20, 80, 0.3, 2)
    g.teile = teile
    return g


def teile(n, laenge, a=True, b=True, zeile=13):
    return [Teil(zeile, 10, "A", i, laenge, a, b, 100, 100) for i in range(1, n + 1)]


# 1) Gemeinsamer Schnitt: 2 × 449 mit geraden Enden passen (20+449+0,3+449+0,3 = 918,6 ≤ 920),
#    mit schrägen Enden nicht (20+449+2,6+449+0,3 = 920,9 > 920)
g = gruppe(teile(2, 449))
r = optimiere(g)
assert len(r) == 1 and gemeinsame_schnitte(r[0]) == 1, r
assert abs(verbrauch(g, r[0]) - 918.6) < 1e-6, verbrauch(g, r[0])
g = gruppe(teile(2, 449, a=False, b=False))
assert len(optimiere(g)) == 2

# 2) Drehen: Teile mit einem schrägen Ende (A schräg). Ohne Drehen liegt schräg an gerade -> 2 Stangen;
#    gedreht liegen die geraden Enden aneinander -> 1 Stange
g = gruppe(teile(2, 449, a=False, b=True))
r = optimiere(g)
assert len(r) == 1 and gemeinsame_schnitte(r[0]) == 1, r
assert len(manuell(g)) == 2

# 3) Gemischt: schräge Enden an Stangenanfang/-ende, gerade in der Mitte
g = gruppe(teile(1, 200, a=False, b=True, zeile=13) + teile(3, 150, zeile=14) + teile(1, 200, a=True, b=False, zeile=15))
r = optimiere(g)
assert len(r) == 1 and gemeinsame_schnitte(r[0]) == 4, [(t.zeile, d) for t, d in r[0]]

# 4) Zufallsaufträge: nie überbelegt, jedes Teil genau einmal, Untergrenze ≤ optimiert ≤ manuell
rnd = random.Random(4711)
for lauf in range(150):
    stange = rnd.choice([3000, 6000, 6500])
    tl = []
    for z in range(rnd.randint(1, 6)):
        L = round(rnd.uniform(40, stange - 101), 1) if rnd.random() < 0.15 else round(rnd.uniform(40, 1900), 1)
        tl += teile(rnd.randint(1, 25), L, a=rnd.random() < 0.8, b=rnd.random() < 0.8, zeile=13 + z)
    g = gruppe(tl, stange)
    r = optimiere(g)
    pruefe([g], [r])
    assert untergrenze(g) <= len(r) <= len(manuell(g)), (lauf, untergrenze(g), len(r), len(manuell(g)))

# 5) Mappe: Beispielauftrag + Sonderfälle
with tempfile.TemporaryDirectory() as d:
    m = Path(d) / "t.xlsm"
    shutil.copy(MAPPE, m)
    wb = load_workbook(m, keep_vba=True)
    wa = wb["Auftrag"]
    zusatz = [  # Zeile 18: zu lang, 19: unbekanntes Profil, 20: schräg/gerade Ende, 21: abweichende Stange
        (60, "Z-zu-lang", "zu lang", "RR 42,4x2,0", "S235-10219", 5950, 1, None, None, None),
        (70, "Z-profil", "Profil fehlt", "RR 99x9", "S235-10219", 500, 2, None, None, None),
        (80, "Z-schraeg", "Gehrung", "RR 42,4x2,0", "S235-10219", 700, 7, 300, 2, None),
        (90, "Z-8m", "lange Stange", "QR 40x40x3", "S235-10219", 2600, 5, 400, 2, 8000),
    ]
    for i, werte in enumerate(zusatz):
        for c, v in enumerate(werte, start=1):
            wa.cell(18 + i, c, v)
    wa.cell(20, 25, "nein"); wa.cell(20, 26, 160.0); wa.cell(20, 27, "ja"); wa.cell(20, 28, 133.2)
    wb["Maschinen"]["F5"] = 8500   # RL-1 darf 8-m-Stangen
    wb.save(m)

    bericht, fehler = optimiere_mappe(m, sicherung=False)
    gr = {f["artikel"]: f["grund"] for f in fehler}
    assert gr == {"Z-zu-lang": "Teil länger als Nutzlänge der Stange", "Z-profil": "Profil unbekannt"}, gr
    st = {(b["profil"], b["stange"]): b for b in bericht}
    assert st[("QR 40x40x3", 8000)]["stangen"] == 2, st[("QR 40x40x3", 8000)]   # 3 × 2600 je 8-m-Stange
    for b in bericht:
        assert b["untergrenze"] <= b["stangen"] <= b["manuell"], b

    wb = load_workbook(m)
    wsp, wdt = wb["Schnittplan"], wb["Stangenbelegung"]
    n_stangen = sum(1 for r in range(5, 1005) if wsp.cell(r, 1).value is not None)
    assert n_stangen == sum(b["stangen"] for b in bericht)
    zeilen = [wdt.cell(r, 5).value for r in range(5, 10005) if wdt.cell(r, 1).value is not None]
    for r in range(13, 22):
        soll = 0 if r in (18, 19) else wb["Auftrag"].cell(r, 7).value
        assert zeilen.count(r) == soll, (r, zeilen.count(r), soll)
    # Gehrungsteile: schräges Ende nie an einem gemeinsamen Schnitt
    for r in range(5, 10005):
        if wdt.cell(r, 5).value == 20:
            # normal: Ende A (schräg) vorn -> eigene Trennung; gedreht: Ende A hinten -> Trennung des Nachfolgers
            pruef = r if wdt.cell(r, 9).value == "normal" else r + 1
            if wdt.cell(pruef, 1).value == wdt.cell(r, 1).value:
                assert wdt.cell(pruef, 10).value != "gemeinsamer Schnitt", r
    # zweiter Lauf überschreibt sauber (keine Reste vom ersten Lauf)
    bericht2, _ = optimiere_mappe(m, sicherung=False)
    wb = load_workbook(m)
    assert sum(1 for r in range(5, 1005) if wb["Schnittplan"].cell(r, 1).value is not None) == n_stangen

print("OK: Optimierung")
