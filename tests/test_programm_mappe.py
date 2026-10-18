"""Vergleicht die Rechnung des Programms mit der Mappe (Abnahme: gleiche Ergebnisse).

Je Testauftrag wird die Mappe mit der Python-Optimierung gefüllt und in LibreOffice durchgerechnet
(alle Formeln: Status, Schnittplan, Kalkulation). Das Programm liest denselben Auftrag und dieselben
Stammdaten aus der Mappe und rechnet selbst. Verglichen werden Schnittplan, Stangenbelegung,
Auswertung, Status und Einplanung je Position sowie die ganze Kalkulation.

Aufruf: python tests/test_programm_mappe.py [Anzahl Zufallsaufträge]   (benötigt LibreOffice)
"""
import random
import shutil
import sys
import tempfile
from pathlib import Path

HIER = Path(__file__).parent
sys.path.insert(0, str(HIER.parent / "programm"))
sys.path.insert(0, str(HIER.parent / "tools"))
sys.path.insert(0, str(HIER))
from openpyxl import load_workbook  # noqa: E402
from optimierung import optimiere_mappe  # noqa: E402
import os  # noqa: E402
import uno  # noqa: E402
from lo_makro import LibreOffice, _pv  # noqa: E402
from test_vba import zufallsauftrag, sonderfaelle, schwer  # noqa: E402
from rohrlaser import rechnen  # noqa: E402
from rohrlaser.mappe_import import lies_mappe, lies_auftrag  # noqa: E402
from rohrlaser.stammdaten import kontext_aus  # noqa: E402

MAPPE = HIER.parent / "mappe" / "Rohrlaser_Kalkulation_v02.xlsm"


def leer(v):
    return v is None or v == ""


def gleich(a, b, tol=1e-6):
    if leer(a) and leer(b):
        return True
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= tol * max(1.0, abs(a), abs(b))
    return a == b


def pruefe(name, was, a, b, tol=1e-6):
    assert gleich(a, b, tol), f"{name} {was}: Programm {a!r}, Mappe {b!r}"


def zusatzwerte(wb, rnd):
    """STEP-Werte, Satz, Gewinde und Preise je Position, damit die ganze Kalkulation durchlaufen wird."""
    wa = wb["Auftrag"]
    for r in range(13, 40):
        if wa.cell(r, 4).value in (None, ""):
            continue
        if wa.cell(r, 1).value in (None, ""):
            # Ohne Pos.-Nr. zeigt Excel in der Kalkulation 0 und rechnet die Zeile, LibreOffice lässt sie leer.
            # Das Programm rechnet wie Excel; für den Vergleich in LibreOffice bekommt jede Zeile eine Nummer.
            wa.cell(r, 1, 900 + r)
        if rnd.random() < 0.7:
            wa.cell(r, 8, round(rnd.uniform(200, 1500), 1))
            wa.cell(r, 9, rnd.randint(0, 8))
            wa.cell(r, 30, round(rnd.uniform(0, 80), 1) if rnd.random() < 0.3 else 0)
            for c in (34, 36, 38):
                wa.cell(r, c, rnd.randint(0, 3))
            wa.cell(r, 35, round(rnd.uniform(0, 60), 1))
            wa.cell(r, 37, round(rnd.uniform(0, 90), 1))
            wa.cell(r, 39, round(rnd.uniform(0, 30), 1))
        if rnd.random() < 0.2:
            wa.cell(r, 32, "2xM6")
            if rnd.random() < 0.6:
                wa.cell(r, 33, rnd.choice(["2xM6", "4xM8; 2xM6", "1xM10"]))
        if rnd.random() < 0.15:
            wa.cell(r, 42, round(rnd.uniform(1, 9), 2))
        if rnd.random() < 0.2:
            wa.cell(r, 43, "Satz A")
            wa.cell(r, 44, rnd.choice([1, 2, 3]))
    for r, z in enumerate(range(5, 15), start=0):
        wb["Gewinde"].cell(z, 3, 10 + r)   # Zeit je Gewinde
    wb["Auftrag"]["C8"] = rnd.choice(["RL-1", "RL-1", "RL-2"])


def mappe_rechnen(quelle, ziel, lo):
    """Python-Optimierung in die Mappe schreiben, dann alle Formeln in LibreOffice rechnen."""
    optimiere_mappe(quelle, sicherung=False)
    doc = lo.desk.loadComponentFromURL(uno.systemPathToFileUrl(os.path.abspath(quelle)), "_blank", 0,
                                       (_pv("Hidden", True),))
    try:
        doc.calculateAll()
        doc.storeToURL(uno.systemPathToFileUrl(os.path.abspath(ziel)), (_pv("FilterName", "Calc MS Excel 2007 XML"),))
    finally:
        doc.close(True)


def vergleichen(name, mappe_eingabe, mappe_werte):
    daten, einst, _ = lies_mappe(mappe_eingabe)
    kopf, positionen, _ = lies_auftrag(mappe_eingabe)
    k = kontext_aus(daten)
    erg = rechnen.optimieren(kopf, positionen, k)
    for p in erg["neue_profile"]:           # wie die Mappe: Profil aus dem Namen anlegen
        daten["profile"].append(p)
    k = kontext_aus(daten)
    wb = load_workbook(mappe_werte, data_only=True)
    wsp, wdt, wau, wa, wk = (wb[n] for n in ("Schnittplan", "Stangenbelegung", "Auswertung", "Auftrag", "Kalkulation"))

    # Schnittplan
    sw = rechnen.stangen_werte(erg, kopf, positionen, k, daten)
    n_mappe = sum(1 for r in range(5, 1005) if wsp.cell(r, 1).value is not None)
    assert n_mappe == len(erg["stangen"]), f"{name}: {len(erg['stangen'])} Stangen gegen {n_mappe}"
    for i, (s, x) in enumerate(zip(erg["stangen"], sw)):
        r = 5 + i
        for c, v in ((1, s["nr"]), (2, s["profil"]), (3, s["werkstoff"]), (4, s["stange"]), (5, s["end"]),
                     (6, s["greif"]), (7, x["teile"]), (8, x["summe"]), (9, s["verlust"]), (10, x["belegt"]),
                     (11, x["rest"]), (12, x["ausnutzung"]), (13, x["pruefung"]), (14, s["trennungen"]),
                     (15, s["gemeinsam"]), (16, s["belegung"]), (17, x["gewicht"]), (18, x["kosten"])):
            pruefe(name, f"Schnittplan Zeile {r} Spalte {c}", v, wsp.cell(r, c).value)
    # Stangenbelegung
    r = 5
    for s in erg["stangen"]:
        for platz, t in enumerate(s["teile"], start=1):
            werte = [s["nr"], platz, t["pos"], t["artikel"], rechnen.ZEILE0 + t["index"], t["laenge"], t["von"],
                     t["bis"], t["lage"], t["trennung"], t["eingespart_mm"] or None, t["eingespart_e"] or None]
            mappe = [wdt.cell(r, c).value for c in range(1, 13)]
            if mappe[3] and str(mappe[3]).startswith("Zeile "):   # Artikel leer: Mappe nennt die Zeile
                werte[3] = mappe[3]
            # Zeilennummer im Programm = Index der gelesenen Positionen; in der Mappe die echte Zeile
            werte[4] = positionen[t["index"]]["zeile"]
            for c, (a, b) in enumerate(zip(werte, mappe), start=1):
                pruefe(name, f"Stangenbelegung Zeile {r} Spalte {c}", a, b)
            r += 1
    assert wdt.cell(r, 1).value is None, f"{name}: Mappe hat mehr Teile"
    # Auswertung
    for i, a in enumerate(erg["auswertung"]):
        rr = 5 + i
        for c, v in ((1, a["profil"]), (2, a["werkstoff"]), (3, a["stange"]), (4, a["teile"]), (5, a["stangen"]),
                     (6, a["manuell"]), (7, a["untergrenze"]), (13, a["ausnutzung_manuell"]), (14, a["gemeinsam"]),
                     (17, a["bewertung"])):
            pruefe(name, f"Auswertung Zeile {rr} Spalte {c}", v, wau.cell(rr, c).value)
    # Status, Einplanung, Positionswerte
    eingeplant = rechnen.eingeplant_je_position(erg)
    for i, p in enumerate(positionen):
        r = p["zeile"]
        w = rechnen.position_werte(p, kopf, k, daten)
        ein = eingeplant.get(rechnen.schluessel(p, i), 0) if erg["stangen"] else None
        st = rechnen.status(p, w, kopf, k, ein)
        pruefe(name, f"Status Zeile {r}", st, wa.cell(r, 23).value)
        if not leer(p.get("profil_id")):
            pruefe(name, f"Eingeplant Zeile {r}", ein or 0, wa.cell(r, 40).value)
            for c, v in ((14, w["kg_m"]), (15, w["gewicht"]), (16, w["laenge_gesamt"]), (24, w["gewicht_position"]),
                         (22, w["vorschub"])):
                pruefe(name, f"Auftrag Zeile {r} Spalte {c}", v, wa.cell(r, c).value)
    # Kalkulation
    einst_k = dict(zuschlag_material=wk["C4"].value, zuschlag_fertigung=wk["C5"].value, gewinn=wk["C6"].value)
    zeilen, summen = rechnen.kalkulation(kopf, positionen, erg, k, daten, einst_k)
    spalten = ["pos", "artikel", "stueck", "material_gruppe", "anteil", "material_teil", "wand_stufe",
               "vorschub_gross", "vorschub_mittel", "vorschub_klein", "vorschub_gravur", "einstechzeit",
               "laser_brutto", "einsparung", "laser_netto", "gewindezeit", "laser_teil", "ruesten_teil",
               "herstell_teil", "preis_teil", "preis_position", "herstell_position", "hinweis",
               "material_position", "laser_position"]
    for i, (p, z) in enumerate(zip(positionen, zeilen)):
        r = p["zeile"]
        if z is None:
            assert leer(wk.cell(r, 2).value), f"{name}: Kalkulation Zeile {r} sollte leer sein"
            continue
        for c, sp in enumerate(spalten, start=1):
            a, b = z[sp], wk.cell(r, c).value
            if sp == "wand_stufe" and a is None:
                a = 0
            if sp in ("pos", "artikel") and a is None:   # Excel zeigt den Bezug auf eine leere Zelle als 0
                a = 0
            pruefe(name, f"Kalkulation Zeile {r} {sp}", a, b, tol=1e-9)
    for c, v in ((4, summen["ruestzeit"]), (5, summen["stangenwechsel"]), (6, summen["stangen"]),
                 (7, summen["ruesten"]), (8, summen["teile"])):
        pruefe(name, f"Kalkulation F{c}", v, wk.cell(c, 6).value, tol=1e-9)
    for c, v in ((4, summen["material"]), (5, summen["laser"]), (6, summen["ruesten"]), (7, summen["herstell"]),
                 (8, summen["verkauf"])):
        pruefe(name, f"Kalkulation V{c}", v, wk.cell(c, 22).value, tol=1e-9)
    return len(erg["stangen"])


def main():
    anzahl = int(sys.argv[1]) if len(sys.argv) > 1 else 12
    tmp = Path(tempfile.mkdtemp())
    rnd = random.Random(20261018)
    faelle = [("Vorlage", None)] + [(f"Zufall {i}", zufallsauftrag) for i in range(anzahl)] + \
             [("Sonderfälle", sonderfaelle), ("Schwer", schwer)]
    with LibreOffice(port=2017) as lo:
        for name, fuellen in faelle:
            m = tmp / "eingabe.xlsm"
            shutil.copy(MAPPE, m)
            if fuellen:
                wb = load_workbook(m, keep_vba=True)
                if fuellen is sonderfaelle:
                    zufallsauftrag(wb, rnd)
                fuellen(wb, rnd)
                zusatzwerte(wb, rnd)
                wb.save(m)
            eingabe = tmp / "nur_eingabe.xlsm"
            shutil.copy(m, eingabe)   # Programm liest den Auftrag vor der Optimierung der Mappe
            mappe_rechnen(m, tmp / "werte.xlsx", lo)
            n = vergleichen(name, eingabe, tmp / "werte.xlsx")
            print(f"{name}: {n} Stangen gleich")
    shutil.rmtree(tmp, ignore_errors=True)
    print("OK: Programm rechnet wie die Mappe")


if __name__ == "__main__":
    main()
