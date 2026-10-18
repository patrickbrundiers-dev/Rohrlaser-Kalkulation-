"""Vergleicht Staffelpreise, Angebot und Werkstattpapier des Programms mit den Makros der Mappe.

Je Testauftrag laufen in LibreOffice die Makros der Mappe (Optimierung.OptimiereTest,
Staffel.StaffelTest, Ausgabe.AusgabeTest); das Programm liest denselben Auftrag samt Staffeln und
rechnet selbst. Verglichen werden jede Zeile des Blatts Staffelpreise, alle Positionen, Summen und
die Preistreppe im Angebot sowie Materialliste und Stangenblöcke im Werkstattpapier.

Aufruf: python tests/test_programm_ausgabe.py [Anzahl Zufallsaufträge]   (benötigt LibreOffice)
"""
import random
import shutil
import sys
import tempfile
from pathlib import Path

HIER = Path(__file__).parent
sys.path.insert(0, str(HIER.parent / "programm"))
sys.path.insert(0, str(HIER.parent / "tools"))
sys.path.insert(0, str(HIER))
from openpyxl import load_workbook  # noqa: E402
from lo_makro import LibreOffice  # noqa: E402
from test_vba import zufallsauftrag  # noqa: E402
from test_programm_mappe import zusatzwerte, pruefe, leer  # noqa: E402
from rohrlaser import ausgabe, rechnen, staffeln  # noqa: E402
from rohrlaser.mappe_import import lies_mappe, lies_auftrag, lies_staffeln  # noqa: E402
from rohrlaser.stammdaten import kontext_aus  # noqa: E402

MAPPE = HIER.parent / "mappe" / "Rohrlaser_Kalkulation_v02.xlsm"


def stufen_eintragen(wb, rnd):
    """Zufällige Mengenstufen: Sätze und Positionen des Auftrags, dazu ein Schlüssel, den es nicht gibt."""
    wa, ws = wb["Auftrag"], wb["Staffeln"]
    for r in range(5, 105):
        for c in (1, 2, 3):
            ws.cell(r, c).value = None
    schl = sorted({str(wa.cell(r, 43).value) for r in range(13, 513) if wa.cell(r, 43).value not in (None, "")})
    schl += [wa.cell(r, 1).value for r in range(13, 40) if wa.cell(r, 1).value not in (None, "")
             and wa.cell(r, 4).value not in (None, "")]
    r = 5
    for s in rnd.sample(schl, min(len(schl), 4)):
        for menge in sorted(rnd.sample([1, 2, 3, 5, 10, 25, 60], rnd.randint(1, 3))):
            ws.cell(r, 1, s)
            ws.cell(r, 2, menge)
            r += 1
    ws.cell(r, 1, "gibt es nicht")
    ws.cell(r, 2, 4)


def programm(eingabe):
    daten, einst, _ = lies_mappe(eingabe)
    kopf, positionen, _ = lies_auftrag(eingabe)
    stufen = lies_staffeln(eingabe)
    k = kontext_aus(daten)
    erg = rechnen.optimieren(kopf, positionen, k)
    daten["profile"] += erg["neue_profile"]          # wie die Mappe: Profil aus dem Namen anlegen
    k = kontext_aus(daten)
    kalk, _ = rechnen.kalkulation(kopf, positionen, erg, k, daten, einst)
    sp = staffeln.rechnen_alle(kopf, positionen, stufen, k, daten, einst)
    ang = ausgabe.angebot(kopf, positionen, kalk, k, einst, sp, dezimal=".")
    wp = ausgabe.werkstatt(kopf, positionen, erg, k, daten)
    return sp, ang, wp


def vergleichen(name, eingabe, werte):
    sp, ang, wp = programm(eingabe)
    wb = load_workbook(werte, data_only=True)

    # Staffelpreise
    wf = wb["Staffelpreise"]
    n = sum(1 for r in range(5, 205) if not leer(wf.cell(r, 1).value))
    assert n == len(sp), f"{name}: {len(sp)} Staffelzeilen gegen {n}"
    for i, z in enumerate(sp):
        for c, f in enumerate(staffeln.SPALTEN, start=1):
            a, b = z[f], wf.cell(5 + i, c).value
            if f == "schluessel":
                a, b = staffeln.txt(a), staffeln.txt(b)
            pruefe(name, f"Staffelpreise Zeile {5 + i} {f}", a, b, tol=1e-9)

    # Angebot: Positionen, Summen, Preistreppe
    wan = wb["Angebot"]
    r = 14
    for z in ang["zeilen"]:
        for c, f in enumerate(("pos", "artikel", "text", "menge", "preis", "gesamt"), start=1):
            a, b = z[f], wan.cell(r, c).value
            if f == "text" and a.endswith("inkl. Material ") and str(b).startswith(a):
                b = a   # Satzteil ohne Werkstoff: Match("") der Mappe trifft eine beliebige Zeile, das Programm lässt leer
            pruefe(name, f"Angebot Zeile {r} {f}", a, b, tol=1e-12)
        r += 1
    assert leer(wan.cell(r, 3).value), f"{name}: Angebot der Mappe hat mehr Zeilen"
    for i, f in enumerate(("netto", "mwst", "brutto")):
        pruefe(name, f"Angebot {f}", ang[f], wan.cell(r + 1 + i, 6).value, tol=1e-12)
    r += 3
    if ang["staffeln"]:
        assert wan.cell(r + 3, 1).value == "Staffelpreise", f"{name}: Preistreppe fehlt in der Mappe"
        r += 5
        for z in ang["staffeln"]:
            for c, f in ((1, "pos"), (3, "text"), (4, "menge"), (5, "preis"), (6, "gesamt")):
                a, b = z[f], wan.cell(r, c).value
                if f == "pos":
                    a, b = staffeln.txt(a), staffeln.txt(b)
                pruefe(name, f"Angebot Preistreppe Zeile {r} {f}", a, b, tol=1e-12)
            r += 1
        assert leer(wan.cell(r, 3).value) and leer(wan.cell(r, 4).value), f"{name}: Mappe hat mehr Staffelzeilen"
    else:
        assert leer(wan.cell(r + 3, 1).value), f"{name}: Mappe hat eine Preistreppe, das Programm nicht"

    # Werkstattpapier
    ww = wb["Werkstattpapier"]
    r = 7
    for m in wp["material"]:
        for c, f in ((1, "profil"), (3, "werkstoff"), (5, "stange"), (6, "stangen"), (7, "gewicht")):
            pruefe(name, f"Materialliste Zeile {r} {f}", m[f], ww.cell(r, c).value, tol=1e-9)
        r += 1
    r += 1
    for m in wp["muster"]:
        assert leer(ww.cell(r - 1, 1).value), f"{name}: Werkstattpapier Zeile {r - 1} sollte leer sein"
        pruefe(name, f"Werkstattpapier Zeile {r} Titel", ausgabe.muster_titel(m), ww.cell(r, 1).value)
        r += 3
        for t in m["teile"]:
            for c, f in enumerate(("platz", "pos", "artikel", "bezeichnung", "laenge", "von", "bis", "lage",
                                   "trennung"), start=1):
                a, b = t[f], ww.cell(r, c).value
                if f == "artikel" and str(b).startswith("Zeile "):   # Artikel leer: Mappe nennt die Blattzeile
                    a = b
                pruefe(name, f"Werkstattpapier Zeile {r} {f}", a, b, tol=1e-9)
            r += 1
        r += 1
    assert all(leer(ww.cell(r + i, 1).value) for i in range(3)), f"{name}: Werkstattpapier der Mappe hat mehr Blöcke"
    return len(sp), len(ang["zeilen"]), len(wp["muster"])


def main():
    anzahl = int(sys.argv[1]) if len(sys.argv) > 1 else 8
    tmp = Path(tempfile.mkdtemp())
    rnd = random.Random(20261019)
    faelle = [("Vorlage", False)] + [(f"Zufall {i}", True) for i in range(anzahl)]
    with LibreOffice(port=2019) as lo:
        for name, zufall in faelle:
            m = tmp / "eingabe.xlsm"
            shutil.copy(MAPPE, m)
            if zufall:
                wb = load_workbook(m, keep_vba=True)
                zufallsauftrag(wb, rnd)
                zusatzwerte(wb, rnd)
                stufen_eintragen(wb, rnd)
                wb.save(m)
            eingabe = tmp / "nur_eingabe.xlsm"
            shutil.copy(m, eingabe)
            lo.ausfuehren(m, [("Optimierung.OptimiereTest", ()), ("Staffel.StaffelTest", ()),
                              ("Ausgabe.AusgabeTest", ())], tmp / "werte.xlsx")
            assert lo.rueckgaben[-1].startswith("OK"), lo.rueckgaben
            n_sp, n_an, n_wp = vergleichen(name, eingabe, tmp / "werte.xlsx")
            print(f"{name}: {n_sp} Staffelstufen, {n_an} Angebotszeilen, {n_wp} Stangenblöcke gleich")
    shutil.rmtree(tmp, ignore_errors=True)
    print("OK: Staffelpreise, Angebot und Werkstattpapier wie in der Mappe")


if __name__ == "__main__":
    main()
