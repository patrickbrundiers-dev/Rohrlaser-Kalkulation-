"""Excel-Export eines Auftrags (Werte, keine Formeln): Auftrag, Kalkulation, Schnittplan, Stangenbelegung,
Auswertung, Staffelpreise und Angebot – zum Weitergeben oder Weiterrechnen in Excel."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .auftrag import KOPF, POSITIONEN
from .auftrag_spalten import AUSW_SPALTEN, KALK_SPALTEN, STANGEN_SPALTEN, TEIL_SPALTEN
from .rechnen import auswertung_zeilen
from .staffeln import SPALTEN as STAFFEL_SPALTEN

KOPF_FUELLUNG = PatternFill("solid", fgColor="1F4E78")
KOPF_SCHRIFT = Font(bold=True, color="FFFFFF")


def _blatt(wb, name, titel, spalten, zeilen, formate=None):
    """spalten: [(Titel, Schlüssel oder Funktion, Breite)]"""
    ws = wb.create_sheet(name)
    ws["A1"] = titel
    ws["A1"].font = Font(bold=True, size=13)
    for c, (t, _, breite) in enumerate(spalten, start=1):
        z = ws.cell(3, c, t)
        z.fill, z.font = KOPF_FUELLUNG, KOPF_SCHRIFT
        z.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = breite
    ws.row_dimensions[3].height = 45
    for r, z in enumerate(zeilen, start=4):
        for c, (_, schl, _) in enumerate(spalten, start=1):
            w = schl(z) if callable(schl) else z.get(schl)
            zelle = ws.cell(r, c, w if isinstance(w, (int, float, str)) or w is None else str(w))
            if formate and spalten[c - 1][0] in formate:
                zelle.number_format = formate[spalten[c - 1][0]]
    ws.freeze_panes = "A4"
    return ws


def auftrag_xlsx(pfad, kopf, positionen, status, kalk_zeilen, summen, ergebnis, stangen_werte, staffelpreise, ang):
    wb = Workbook()
    wb.remove(wb.active)
    # Auftrag: Kopf und Positionen
    ws = _blatt(wb, "Auftrag", f"Auftrag {kopf.get('nummer') or ''}",
                [(f.titel, f.name, max(8, f.breite / 7)) for f in POSITIONEN.felder] + [("Status", "_status", 40)],
                [dict(p, _status=s) for p, s in zip(positionen, status)])
    ws.insert_rows(2, len(KOPF) + 1)
    for i, f in enumerate(KOPF):
        ws.cell(2 + i, 1, f.titel).font = Font(bold=True)
        ws.cell(2 + i, 3, kopf.get(f.name))
    ws.freeze_panes = f"A{len(KOPF) + 5}"
    _blatt(wb, "Kalkulation", "Kalkulation", [(t, s, w / 6) for t, s, _, w in KALK_SPALTEN], [z for z in kalk_zeilen if z])
    ks = wb["Kalkulation"]
    ks.insert_rows(2, 6)
    for i, (t, n) in enumerate((("Material [€]", "material"), ("Laser + Gewinde [€]", "laser"),
                                ("Rüsten + Stangenwechsel [€]", "ruesten"), ("Herstellkosten [€]", "herstell"),
                                ("Verkaufspreis Auftrag [€]", "verkauf"))):
        ks.cell(2 + i, 1, t).font = Font(bold=True)
        ks.cell(2 + i, 3, summen[n]).number_format = "#,##0.00"
    if ergebnis:
        stangen = [dict(s, **w) for s, w in zip(ergebnis["stangen"], stangen_werte)]
        _blatt(wb, "Schnittplan", "Schnittplan", [(t, s, w / 6) for t, s, _, w in STANGEN_SPALTEN], stangen)
        teile = [dict(t, nr=s["nr"], platz=i) for s in ergebnis["stangen"] for i, t in enumerate(s["teile"], start=1)]
        _blatt(wb, "Stangenbelegung", "Stangenbelegung",
               [("Stange Nr.", "nr", 9)] + [(t, s, w / 6) for t, s, _, w in TEIL_SPALTEN], teile)
        _blatt(wb, "Auswertung", "Auswertung", [(t, s, w / 6) for t, s, _, w in AUSW_SPALTEN],
               auswertung_zeilen(ergebnis, stangen_werte))
        if ergebnis["fehler"]:
            _blatt(wb, "Nicht eingeplant", "Nicht eingeplante Teile",
                   [("Pos.", "pos", 8), ("Artikelnummer", "artikel", 20), ("Anzahl", "anzahl", 8), ("Grund", "grund", 50)],
                   ergebnis["fehler"])
    titel = ["Satz oder Pos.", "Bezeichnung", "Einheit", "Menge", "Teile gesamt [Stk]", "Stangen", "Material [€]",
             "Laser + Gewinde [€]", "Rüsten + Stangenwechsel [€]", "Herstellkosten [€]", "Herstellkosten je Einheit [€]",
             "Verkaufspreis je Einheit [€]", "Verkaufspreis gesamt [€]", "Hinweis"]
    _blatt(wb, "Staffelpreise", "Staffelpreise", [(t, s, 14) for t, s in zip(titel, STAFFEL_SPALTEN)],
           staffelpreise or [])
    _blatt(wb, "Angebot", "Angebot (Positionen)",
           [("Pos.", "pos", 7), ("Artikelnummer", "artikel", 16), ("Bezeichnung", "text", 70), ("Menge [Stk]", "menge", 10),
            ("Stückpreis [€]", "preis", 12), ("Gesamtpreis [€]", "gesamt", 13)],
           ang["zeilen"] + [{}, dict(text="Summe netto [€]", gesamt=ang["netto"]),
                            dict(text="zzgl. MwSt. [€]", gesamt=ang["mwst"]),
                            dict(text="Gesamtbetrag [€]", gesamt=ang["brutto"])])
    wb.save(pfad)
