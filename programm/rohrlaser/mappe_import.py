"""Stammdaten und Einstellungen aus einer Kalkulationsmappe (v02, .xlsm/.xlsx) lesen.

Gelesen werden nur Eingabezellen: die Stammdatenblätter ab Zeile 5, die Zuschläge der Kalkulation,
Angebotskopf und Auftragskopf (Maschine, Stangenlänge, Auftragsnummer). Funktioniert auch mit
abgelegten Auftragsmappen, weil sie dieselben Blätter haben.
"""
import datetime
from pathlib import Path

from openpyxl import load_workbook

from .stammdaten import TABELLEN, EINSTELLUNGEN, zahl_aus_text, zahl_als_text


def _wert(f, v):
    """Zellwert in den gespeicherten Wert umwandeln (Zahlen auch als Text mit Komma)."""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if f.art in ("zahl", "ganz"):
        if isinstance(v, bool):
            raise ValueError
        w = float(v) if isinstance(v, (int, float)) else zahl_aus_text(v)
        return int(w) if f.art == "ganz" and w == int(w) else w
    if f.art == "datum":
        if isinstance(v, datetime.datetime):
            return v.date().isoformat()
        if isinstance(v, datetime.date):
            return v.isoformat()
        return datetime.datetime.strptime(str(v).strip(), "%d.%m.%Y").date().isoformat()
    if isinstance(v, float):
        return f"{v:g}"
    return str(v).strip()


def lies_mappe(pfad):
    """Ergebnis: (daten {Tabelle: [Zeile]}, einstellungen {Name: Wert}, hinweise [Text]).
    Zeilen mit doppeltem Schlüssel werden übersprungen und als Hinweis gemeldet; unlesbare Werte
    bleiben leer und werden gemeldet."""
    pfad = Path(pfad)
    wb = load_workbook(pfad, read_only=True, data_only=True, keep_vba=False)
    fehlend = [t.blatt for t in TABELLEN if t.blatt not in wb.sheetnames]
    if fehlend:
        wb.close()
        raise ValueError(f"'{pfad.name}' ist keine Kalkulationsmappe: Blatt {', '.join(fehlend)} fehlt.")
    daten, hinweise = {}, []
    for t in TABELLEN:
        ws = wb[t.blatt]
        zeilen, schluessel = [], set()
        breite = max(f.spalte for f in t.felder)
        for nr, werte in enumerate(ws.iter_rows(min_row=t.zeilen[0], max_row=t.zeilen[1], max_col=breite,
                                                values_only=True), start=t.zeilen[0]):
            werte = list(werte) + [None] * (breite - len(werte))
            if all(werte[f.spalte - 1] in (None, "") for f in t.felder):
                continue
            z = {}
            for f in t.felder:
                roh = werte[f.spalte - 1]
                try:
                    z[f.name] = _wert(f, roh)
                except (ValueError, TypeError):
                    z[f.name] = None
                    hinweise.append(f"{t.blatt} Zeile {nr}: {f.titel} '{roh}' nicht lesbar, bleibt leer.")
            schl = tuple(z.get(s) for s in t.schluessel)
            if all(s is not None for s in schl):
                if schl in schluessel:
                    text = " + ".join(str(s) if not isinstance(s, float) else zahl_als_text(s) for s in schl)
                    hinweise.append(f"{t.blatt} Zeile {nr}: '{text}' doppelt, nicht übernommen.")
                    continue
                schluessel.add(schl)
            zeilen.append(z)
        daten[t.name] = zeilen

    einstellungen = {}
    for e in EINSTELLUNGEN:
        if e.zelle is None or e.zelle[0] not in wb.sheetnames:
            continue
        v = wb[e.zelle[0]][e.zelle[1]].value
        if v is None or (isinstance(v, str) and v.startswith("=")):
            continue
        try:
            if e.art in ("zahl", "anteil"):
                v = float(v)
            elif e.art == "ganz":
                v = int(float(v))
            else:
                v = str(v).strip()
        except (TypeError, ValueError):
            hinweise.append(f"{e.zelle[0]} {e.zelle[1]}: {e.titel} '{v}' nicht lesbar, Einstellung bleibt.")
            continue
        einstellungen[e.name] = v
    if "Listen" in wb.sheetnames:
        formen = [r[0] for r in wb["Listen"].iter_rows(min_row=5, max_row=20, max_col=1, values_only=True)
                  if r[0] not in (None, "")]
        if formen:
            einstellungen["profilformen"] = "\n".join(str(f).strip() for f in formen)
    wb.close()
    return daten, einstellungen, hinweise
